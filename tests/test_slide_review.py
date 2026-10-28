import json
import sqlite3
from contextlib import closing

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Inches

from battery_insight import slide_link as sl
from battery_insight import slide_review as sr
from battery_insight.slide_context import COLS, data_only_context


def _table_deck(path, rows, title="결과"):
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = title
    tbl = s.shapes.add_table(len(rows), len(rows[0]), Inches(0.5), Inches(1.5), Inches(8), Inches(3)).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            tbl.cell(i, j).text = str(v)
    prs.save(str(path))


def _project(tmp_path):
    root = tmp_path / "2팀" / "홍길동"
    root.mkdir(parents=True)
    rows = [["Capacity (mAh/g)", "1st Charge", "1st Discharge", "Eff (%)"],
            ["Set1: PTFE1.5_ZQ81 (7/7)", "225.0", "202.5", "90.0"],
            ["Set2: PTFE1.5_ZQ82 (7/7)", "225.0", "196.0", "87.1"]]
    _table_deck(root / "260720 ZQ 평가_v1.pptx", rows, title="ZQ81 ZQ82 초기 용량 결과")
    _table_deck(root / "260725 ZQ 평가_v2.pptx", rows, title="ZQ81 ZQ82 초기 용량 결과")   # same slide, newer version
    proj = tmp_path / "proj"
    (proj / "dx").mkdir(parents=True)
    (proj / "db").mkdir()
    (proj / "slides").mkdir()
    pd.DataFrame([
        dict(test_key="a1", owner_name="홍길동", requester_code="00123456", test_date="2026-07-07", set_no="1",
             sample_name="PTFE1.5 ZQ81", test_name="x_260707_Set1_ZQ81", n_channels=2),
        dict(test_key="b2", owner_name=None, requester_code="00999999", test_date="2026-07-07", set_no="2",
             sample_name="PTFE1.5 ZQ82", test_name="x_260707_Set2_ZQ82", n_channels=2),
        dict(test_key="c3", owner_name="홍길동", requester_code="00123456", test_date="2026-07-08", set_no="3",
             sample_name="QQ77", test_name="x_260708_Set3_QQ77", n_channels=1),
        dict(test_key="d4", owner_name="홍길동", requester_code="00123456", test_date="2026-07-07", set_no="1",
             sample_name="PTFE1.5 ZQ81", test_name="y_260707_Set1_ZQ81", n_channels=1),
    ]).to_csv(proj / "dx" / "dx_tests.csv", index=False)
    cyc = []
    for key, ice in (("a1", 0.900), ("b2", 0.871), ("d4", 0.900)):
        for n in range(1, 4):
            cyc.append(dict(run_id=f"{key}:ch1", cycle_index=n, discharge_capacity_ah=0.0045 * (ice if n == 1 else 1),
                            charge_capacity_ah=0.0045 * (1 if n == 1 else 1.001)))
    pd.DataFrame(cyc).to_parquet(proj / "db" / "run_cycles.parquet")
    c = sqlite3.connect(proj / "db" / "research.sqlite3")
    c.execute("create table tests(test_key, owner, test_date, set_no, sample_name, test_type, design_id, candidate_id)")
    c.execute("create table designs(design_id, design_label, status)")
    c.executemany("insert into tests(test_key, design_id) values (?, ?)", [("a1", "D1"), ("b2", "D2"), ("c3", "D3")])
    c.executemany("insert into designs values (?, ?, 'confirmed')", [("D1", "ZQ81"), ("D2", "ZQ82"), ("D3", "QQ77")])
    c.commit()
    c.close()
    return tmp_path / "2팀", proj


def test_review_overrides_links_and_writes_slide_facts(tmp_path):
    root, proj = _project(tmp_path)
    pd.DataFrame([
        dict(test_key="b2", decision="link", file_name="260720 ZQ 평가_v1.pptx", slide_no="1", set_label="Set2",
             confidence="high", evidence="ICE 87.1", slide_design="PTFE1.5 ZQ82", slide_conditions="25℃",
             owner_guess="임꺽정", reviewer="agent"),
        dict(test_key="c3", decision="no_slide", note="보고 전", reviewer="agent"),
        dict(test_key="d4", decision="link", file_name="없는 덱.pptx", slide_no="3", set_label="Set1",
             confidence="high", reviewer="agent"),
    ]).to_csv(proj / "slides" / sr.REVIEW_FILE, index=False, encoding="utf-8-sig")

    res = sl.build_slide_link(proj, [root], workers=1, log=lambda *_: None)
    links = res["links"]
    b2 = links[links["test_key"] == "b2"]
    assert list(b2["confidence"]) == ["confirmed"], b2.to_dict("records")
    assert b2.iloc[0]["file_name"] == "260720 ZQ 평가_v1.pptx" and b2.iloc[0]["review_design"] == "PTFE1.5 ZQ82"
    assert "Set2" in b2.iloc[0]["set_metrics"] or "87.1" in b2.iloc[0]["set_metrics"]
    assert "c3" not in set(links["test_key"])
    tb = res["tests_without_slides"].set_index("test_key")
    assert tb.loc["c3", "review_status"] == "no_slide"
    assert tb.loc["c3", "ctx_reason"] == "사이클 데이터 없음"
    assert "b2" not in set(res["context"]["test_key"])
    # unresolved review row: reported, and the test keeps its automatic links
    assert list(res["review_unresolved"]["test_key"]) == ["d4"]
    assert not res["review"].set_index("test_key").loc["d4", "resolved"]
    assert (links.loc[links["test_key"] == "d4", "review_status"] == "unresolved").all()
    assert len(links[links["test_key"] == "d4"])

    f = res["test_facts"].set_index("test_key")
    assert f.loc["b2", "link_status"] == "confirmed" and f.loc["b2", "review_status"] == "link"
    assert f.loc["b2", "slide_ice_pct"] == 87.1
    assert abs(f.loc["b2", "db_ice_pct"] - 87.1) < 0.01
    assert abs(f.loc["b2", "implied_active_mass_mg"] - 1000 * 4.5 * 0.871 / 196.0) < 0.01
    assert f.loc["b2", "slide_design"] == "PTFE1.5 ZQ82"
    assert f.loc["c3", "link_status"] == "no_slide"
    assert f.loc["d4", "review_status"] == "unresolved"

    own = res["owner_candidates"].set_index("candidate_name")
    assert set(own["requester_code"]) == {"00999999"}               # code kept as text (leading zeros)
    assert own.loc["임꺽정", "from_review"] == 1 and own.loc["홍길동", "from_decks"] == 1

    c = sqlite3.connect(proj / "db" / "research.sqlite3")
    try:
        row = c.execute("select link_status, slide_file_name, design_label from v_test_slide_facts "
                        "where test_key='b2'").fetchone()
        assert row == ("confirmed", "260720 ZQ 평가_v1.pptx", "ZQ82")
        d2 = c.execute("select n_tests_with_slides, n_confirmed, slide_designs from design_slide_facts "
                       "where design_id='D2'").fetchone()
        assert d2[0] == 1 and d2[1] == 1 and "ZQ82" in d2[2]
        assert "D2" in set(res["design_facts"]["design_id"])
        assert c.execute("select count(*) from slide_link_review").fetchone()[0] == 3
        source = c.execute("select status, reviewer_kind, human_confirmed from evidence_sources "
                           "where test_key='b2'").fetchone()
        assert source == ("eligible", "agent", 0)
        inferred = c.execute("select proposed_value, status from evidence_attributes "
                             "where test_key='b2' and field='temperature_c'").fetchone()
        assert inferred == (25.0, "candidate")
        assert c.execute("select review_status from v_test_slides where test_key='b2'").fetchone()[0] == "link"
        assert c.execute("select ctx_reason from v_test_slide_facts where test_key='c3'").fetchone()[0] == "사이클 데이터 없음"
        assert c.execute("select ctx_reason from v_test_slide_facts where test_key='b2'").fetchone()[0] is None
        fam = c.execute("select family_size, family_rank, file_name from slide_decks order by family_rank").fetchall()
        assert [r[:2] for r in fam] == [(2, 1), (2, 2)] and fam[0][2] == "260725 ZQ 평가_v2.pptx"
    finally:
        c.close()

    # the two decks repeat one slide: only the newest copy is canonical
    sets = pd.read_parquet(proj / "slides" / "slide_sets.parquet")
    assert sets["content_hash"].nunique() == 1
    assert set(sets.loc[sets["is_canonical"], "file_name"]) == {"260725 ZQ 평가_v2.pptx"}


def test_family_key_groups_versions():
    k = sr.family_key
    assert k("260911 Si계 건식음극 전해질함량 DOE_v4.2.pptx") == k("260923 Si계 건식음극 전해질함량 DOE_v5.pptx")
    assert k("260902 (update) 건조 task 정리.pptx") == k("260830 건조 task 정리.pptx")
    assert k("261007 반제품 드라이룸 노출 평가_v4.1.pptx") == k("260930 반제품 드라이룸 노출 평가_v3_이수현.pptx")
    assert k("260519 SJ1300B 평가.pptx") != k("260519 SJ1300A 평가.pptx")


def test_human_review_wins_and_cached_refresh_uses_current_qc(tmp_path):
    root, proj = _project(tmp_path)
    review = [
        dict(test_key="a1", decision="link", file_name="260720 ZQ 평가_v1.pptx", slide_no="1",
             set_label="Set1", confidence="high", reviewer="agent-2"),
        dict(test_key="b2", decision="link", file_name="260720 ZQ 평가_v1.pptx", slide_no="1",
             set_label="Set2", confidence="high", reviewer="agent"),
        dict(test_key="a1", decision="no_slide", note="not this experiment", reviewer="researcher"),
    ]
    pd.DataFrame(review).to_csv(proj / "slides" / sr.REVIEW_FILE, index=False)
    first = sl.build_slide_link(proj, [root], log=lambda *_: None)
    assert "a1" not in set(first["links"]["test_key"])
    assert sr.review_status(pd.DataFrame(review))["a1"] == "no_slide"
    assert first["test_facts"].set_index("test_key").loc["b2", "db_ice_pct"] > 80
    pd.DataFrame([dict(run_id="b2:ch1", qc_status="defect", qc_reasons="excluded")]).to_parquet(
        proj / "db" / "run_qc.parquet")
    second = sl.build_slide_link(proj, log=lambda *_: None)
    assert pd.isna(second["test_facts"].set_index("test_key").loc["b2", "db_ice_pct"])
    assert second["test_facts"].set_index("test_key").loc["b2", "db_cycle_basis"] == "no_eligible_cycles"
    with sqlite3.connect(proj / "db" / "research.sqlite3") as conn:
        assert conn.execute("SELECT db_ice_pct FROM test_slide_facts WHERE test_key='b2'").fetchone() == (None,)
    excel = pd.read_excel(proj / "db" / "research.xlsx", sheet_name="test_slide_facts")
    assert pd.isna(excel.set_index("test_key").loc["b2", "db_ice_pct"])


def test_summary_values_belong_to_primary_source_and_distinct_sets_survive(tmp_path):
    root, proj = _project(tmp_path)
    pd.DataFrame([
        dict(test_key="a1", decision="link", file_name="260720 ZQ 평가_v1.pptx", slide_no="1",
             set_label=label, confidence="high", reviewer="agent")
        for label in ("Set1", "Set2")
    ]).to_csv(proj / "slides" / sr.REVIEW_FILE, index=False)
    result = sl.build_slide_link(proj, [root], log=lambda *_: None)
    links = result["links"]
    assert len(links[links["test_key"] == "a1"]) == 2
    links = links[links["test_key"] == "a1"].copy()
    links["score"] = [30, 25]
    links.iloc[0, links.columns.get_loc("set_metrics")] = '{"ice_pct":90}'
    links.iloc[1, links.columns.get_loc("set_metrics")] = '{"discharge_1st_mah_g":123}'
    facts = sr.build_test_facts(links, sr.load_review(proj), sl._load_tests(proj), sl.test_cycle_facts(proj))
    row = facts.set_index("test_key").loc["a1"]
    assert json.loads(row["slide_metrics_json"]) == {"ice_pct": 90}
    assert pd.isna(row["implied_active_mass_mg"])


def test_duplicate_deck_names_need_explicit_path(tmp_path):
    root, proj = _project(tmp_path)
    source = root / "홍길동" / "260720 ZQ 평가_v1.pptx"
    duplicate = root / "임꺽정" / source.name
    duplicate.parent.mkdir()
    duplicate.write_bytes(source.read_bytes())
    review = dict(test_key="a1", decision="link", file_name=source.name, slide_no="1", set_label="Set1",
                  confidence="high", reviewer="human", evidence="Review of exact source Set")
    pd.DataFrame([review]).to_csv(proj / "slides" / sr.REVIEW_FILE, index=False)
    result = sl.build_slide_link(proj, [root], log=lambda *_: None)
    assert list(result["review_unresolved"]["test_key"]) == ["a1"]
    review["file_path"] = str(source)
    pd.DataFrame([review]).to_csv(proj / "slides" / sr.REVIEW_FILE, index=False)
    resolved = sl.build_slide_link(proj, log=lambda *_: None)
    assert resolved["review_unresolved"].empty
    assert list(resolved["links"].query("test_key == 'a1'")["file_path"]) == [str(source)]
    with sqlite3.connect(proj / "db" / "research.sqlite3") as conn:
        assert conn.execute("SELECT status FROM evidence_sources WHERE test_key='a1'").fetchone() == ("eligible",)


def test_multi_set_label_does_not_inherit_first_set():
    sets = pd.DataFrame([dict(set_no=1, set_label="Set1: A"), dict(set_no=3, set_label="Set3: B")])
    assert sr._sets_for_label("Set1-3", [0, 1], sets) == []
    assert sr._sets_for_label("Set1 and Set3", [0, 1], sets) == []
    assert sr._sets_for_label("Set#1", [0, 1], sets) == [0]
    assert sr._sets_for_label("Set1_ref", [0, 1], sets) == [0]


@pytest.mark.parametrize("missing", ["evidence", "reviewer"])
def test_empty_raw_review_metadata_does_not_create_inference(tmp_path, missing):
    root, proj = _project(tmp_path)
    row = dict(test_key="a1", decision="link", file_name="260720 ZQ 평가_v1.pptx", slide_no="1",
               set_label="Set1", confidence="high", reviewer="agent", evidence="Checked exact Set",
               slide_conditions="25℃")
    row[missing] = ""
    pd.DataFrame([row]).to_csv(proj / "slides" / sr.REVIEW_FILE, index=False)
    sl.build_slide_link(proj, [root], log=lambda *_: None)
    with closing(sqlite3.connect(proj / "db" / "research.sqlite3")) as conn:
        assert conn.execute("SELECT status FROM evidence_sources WHERE test_key='a1'").fetchone() == ("rejected",)
        assert conn.execute("SELECT COUNT(*) FROM evidence_attributes WHERE test_key='a1'").fetchone()[0] == 0


def test_review_status_priority_and_load(tmp_path):
    (tmp_path / "slides").mkdir()
    pd.DataFrame([dict(test_key="t", decision="unsure"), dict(test_key="t", decision="LINK", file_name="x", slide_no="1"),
                  dict(test_key="u", decision="no_slide"), dict(test_key="v", decision="maybe")]
                 ).to_csv(tmp_path / "slides" / sr.REVIEW_FILE, index=False)
    r = sr.load_review(tmp_path)
    assert set(r["test_key"]) == {"t", "u"}             # unknown decisions are dropped
    st = sr.review_status(r)
    assert st["t"] == "link" and st["u"] == "no_slide"


def test_data_only_context_reasons():
    def t(key, date, owner="갑", sample="Set1", **kw):
        return dict(test_key=key, owner_base=owner, test_date=date, sample_name=sample, test_name=f"x_{sample}", **kw)
    tests = pd.DataFrame([
        t("t1", "2026-05-01", tags="contract_batch:618"),                                   # reported
        t("t2", "2026-05-01", sample="Set2 FM", test_type_final="Formation", tags="contract_batch:618"),
        t("t3", "2025-10-01"),
        t("t4", "2026-05-28"),
        t("t5", "2026-05-01", owner="을"),
        t("t6", "2026-05-01", sample="충방전기 비교 500mA vs 6A"),
        t("t7", "2026-04-15"),
        t("t8", "2026-04-16", sample="In situ XRD"),
        t("t9", "2026-04-20"),
        t("t10", "2026-05-10", owner=None, requester_code="07000664"),
        *[t(k, "2026-06-10", owner="병") for k in ("t11", "t12", "t13")],
        t("u1", "2026-07-01", owner="병", sample="건식음극 집전체평가 Ref 1um"),                # reported
        t("u2", "2026-07-01", owner="병", sample="건식음극 전해질함량평가 SJ900 8대2"),
        t("u3", "2026-07-01", owner="병", sample="건식음극 전해질함량평가 SJ900 6대4 re"),
        t("u4", "2026-07-01", owner="병", sample="건식음극 집전체평가 Ref. 2um"),
    ])
    facts = pd.DataFrame(dict(test_key=tests["test_key"], link_status="no_slide", review_status="", review_note="",
                              slide_file_name="", slide_no=None, slide_set_label="", slide_design="", db_ice_pct=90.0))
    facts.loc[0, ["link_status", "slide_file_name", "slide_no", "slide_set_label"]] = ["confirmed", "260510 평가.pptx", 3, "Set1"]
    facts.loc[facts["test_key"] == "u1", ["link_status", "slide_file_name", "slide_no"]] = ["confirmed", "260716 집전체.pptx", 2]
    facts.loc[facts["test_key"] == "t7", "db_ice_pct"] = 99.5
    decks = pd.DataFrame(dict(slide_owner=["갑", "갑", "병", "병"],
                              file_date=["2026-03-01", "2026-06-01", "2026-01-01", "2026-12-01"]))
    owners = pd.DataFrame([dict(requester_code="07000664", candidate_name="갑", suggestion="유력")])
    design = {"t1": "D1", "t7": "D1", "t9": "D9"}
    ncyc = {k: 100 for k in tests["test_key"]} | {"t2": 3, "t9": 0}

    ctx = data_only_context(tests, facts, decks, owners, design, ncyc)
    assert list(ctx.columns) == COLS
    x = ctx.set_index("test_key")
    assert "t1" not in x.index and "u1" not in x.index and len(x) == 15
    assert x.loc["t2", "ctx_reason"] == "선행·점검 단계(본시험 장표 있음)"
    assert x.loc["t2", "group_basis"] == "같은 도급 배치" and x.loc["t2", "sibling_slide"] == "260510 평가.pptx p3 Set1"
    assert x.loc["t3", "ctx_reason"] == "장표 보관 이전 시험"
    assert x.loc["t4", "ctx_reason"] == "미보고(최신 장표 이후)"
    assert x.loc["t5", "ctx_reason"] == "담당자 장표 없음"
    assert x.loc["t6", "ctx_reason"] == "장비·채널 점검"
    assert x.loc["t7", "ctx_reason"] == "같은 실험군은 보고됨(이 시험 미표시)" and x.loc["t7", "group_basis"] == "같은 설계"
    assert "화성 후 이어서" in x.loc["t7", "ctx_flags"] and "활성화 뒤" in x.loc["t7", "ctx_detail"]
    assert x.loc["t8", "ctx_reason"] == "분석용 셀(XRD/DSC 등)"
    assert x.loc["t9", "ctx_reason"] == "사이클 데이터 없음" and "사이클 데이터 없음" in x.loc["t9", "ctx_flags"]
    assert (x.loc["t10", "owner_eff"], x.loc["t10", "owner_basis"], x.loc["t10", "ctx_reason"]) == ("갑", "사번 후보", "미상")
    assert set(x.loc[["t11", "t12", "t13"], "ctx_reason"]) == {"실험군 전체 미보고"}
    assert x.loc["t11", "n_group"] == 3 and x.loc["t11", "group_basis"] == "같은 투입일"
    # one input day, two sub-experiments: only 집전체 was reported
    assert x.loc["u4", "ctx_reason"] == "같은 실험군은 보고됨(이 시험 미표시)"
    assert x.loc["u4", "sibling_slide"] == "260716 집전체.pptx p2" and x.loc["u4", "sibling_common"] == "ref,집전체평가"
    assert set(x.loc[["u2", "u3"], "ctx_reason"]) == {"같은 투입분 다른 실험만 보고됨"}
    assert x.loc["u2", "sibling_slide"] == "" and x.loc["u2", "n_group"] == 2
    assert "다른 실험 1건만 보고: 260716 집전체.pptx p2" in x.loc["u2", "ctx_detail"]
