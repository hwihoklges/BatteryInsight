"""The user-facing append-one-document workflow preserves input and provenance."""
import csv
import json
from pathlib import Path

import pytest

from battery_insight.cli import main
from battery_insight.knowledge_registry import KnowledgeRegistry


def invoke(*args):
    return main([str(arg) for arg in args])


def save_rows(path, rows):
    path.write_text(json.dumps({"records": rows}, ensure_ascii=False), encoding="utf-8-sig")
    return path


def make_pptx(path, text="2026-01-29 SET 001\nSample A"):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches
    path.parent.mkdir(parents=True, exist_ok=True)
    document = pptx.Presentation()
    slide = document.slides.add_slide(document.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(4), Inches(2)).text = text
    document.save(path)
    return path


@pytest.fixture
def workspace(tmp_path):
    root = tmp_path / "knowledge"
    assert invoke("knowledge-init", "--out", root) == 0
    return root, root / "knowledge.sqlite3"


def snapshot(path):
    with KnowledgeRegistry(path, read_only=True) as registry:
        return registry.snapshot()


def intake(workspace, tmp_path, text="2026-01-29 SET 001\nSample A"):
    root, registry = workspace
    source = make_pptx(tmp_path / "input" / "results.pptx", text)
    assert invoke("knowledge-intake", "--registry", registry, "--documents", source, "--out", root / "review") == 0
    evidence_id = snapshot(registry)["evidence"][0]["evidence_id"]
    return source, evidence_id


def reviewed_row(evidence_id, **changes):
    return {"record_id": "rec_A", "sample_id": "sample_A", "sample_name": "Sample A",
            "experiment_name": "Pressure experiment", "test_date": "2026-01-29", "date_kind": "test_start",
            "namespace": "07100395", "set_no": "001", "test_type": "RPT", "rpt_stage": "BOL",
            "status": "confirmed", "evidence_ids": [evidence_id]} | changes


def records_command(registry, path, *, reviewed=False, out=None):
    args = ["knowledge-records", "--registry", registry, "--records", path]
    if reviewed:
        args += ["--reviewer", "Operator", "--reason", "Checked original slide row"]
    if out:
        args += ["--out", out]
    return invoke(*args)


def test_init_is_idempotent_and_preserves_user_edited_templates(workspace):
    root, registry = workspace
    edits = {"records.template.json": '{"records": [{"sample_name": "User entry"}]}',
             "decisions.template.json": '{"decisions": [{"note": "preserve"}]}',
             "raw_keys.template.csv": "experiment_id,set_no\nraw-1,0007\n",
             "required_information.json": '{"custom_note":"kept"}'}
    for name, value in edits.items():
        (root / name).write_text(value, encoding="utf-8")
    before = snapshot(registry)
    assert invoke("knowledge-init", "--out", root) == 0
    assert snapshot(registry) == before
    for name, value in edits.items():
        assert (root / name).read_text(encoding="utf-8") == value
    assert (root / "index.html").exists()


def test_intake_is_idempotent_and_never_creates_sample_facts(workspace, tmp_path):
    root, registry = workspace
    source, _ = intake(workspace, tmp_path)
    source_before, before = source.read_bytes(), snapshot(registry)
    assert invoke("knowledge-intake", "--registry", registry, "--documents", source, "--out", root / "review") == 0
    assert snapshot(registry) == before
    assert source.read_bytes() == source_before
    assert before["records"] == [] and before["samples"] == []
    assert len(before["documents"]) == 1
    assert all(row["needs_review"] for row in before["evidence"])


def test_draft_confirm_update_and_export_preserve_ids_revisions_and_dates(workspace, tmp_path):
    root, registry = workspace
    _, evidence_id = intake(workspace, tmp_path)
    path = save_rows(tmp_path / "records.json", [reviewed_row(evidence_id, status="draft")])
    assert records_command(registry, path, out=root / "review") == 0
    first = snapshot(registry)["records"][0]
    assert first["revision"] == 1 and not first["matching_eligible"]
    save_rows(path, [{"record_id": "rec_A", "revision": 1, "status": "confirmed", "report_date": "2026-02-04"}])
    assert records_command(registry, path, reviewed=True) == 0
    second = snapshot(registry)["records"][0]
    assert second["revision"] == 2 and second["matching_eligible"]
    save_rows(path, [{"record_id": "rec_A", "revision": 2, "notes": "Date clarified"}])
    assert records_command(registry, path, reviewed=True) == 0
    assert invoke("knowledge-export", "--registry", registry, "--out", root / "export") == 0
    export = json.loads((root / "export" / "knowledge_snapshot.json").read_text(encoding="utf-8"))
    row = export["records"][0]
    assert row["record_id"] == "rec_A" and row["revision"] == 3
    assert row["set_no"] == "001" and row["test_date"] == "2026-01-29" and row["report_date"] == "2026-02-04"
    assert row["physical_cell_id"] is None
    with (root / "export" / "records.csv").open(encoding="utf-8-sig", newline="") as stream:
        assert list(csv.DictReader(stream))[0]["set_no"] == "001"
    with KnowledgeRegistry(registry, read_only=True) as reader:
        assert [r["revision"] for r in reader.list_revisions("rec_A")] == [1, 2, 3]


def test_confirm_requires_review_and_error_preserves_draft(workspace, tmp_path):
    _, registry = workspace
    _, evidence_id = intake(workspace, tmp_path)
    path = save_rows(tmp_path / "records.json", [reviewed_row(evidence_id, status="draft")])
    assert records_command(registry, path) == 0
    before = snapshot(registry)
    save_rows(path, [{"record_id": "rec_A", "status": "confirmed"}])
    assert records_command(registry, path) == 2
    assert snapshot(registry) == before


@pytest.mark.parametrize("payload", ["{broken", '{"records":null}', '{"records":["text"]}',
                                     '{"records":[{"set_no":1}]}',
                                     '{"records":[{"test_date":"26-01-29"}]}',
                                     '{"records":[{"metadata":{"raw_path":"overwrite"}}]}'])
def test_invalid_record_inputs_are_controlled_errors_and_preserve_registry(workspace, tmp_path, payload, capsys):
    _, registry = workspace
    before = snapshot(registry)
    path = tmp_path / "bad-records.json"
    path.write_text(payload, encoding="utf-8")
    assert records_command(registry, path) == 2
    assert "ERROR:" in capsys.readouterr().err
    assert snapshot(registry) == before


def test_export_escapes_document_record_and_spreadsheet_formula_text(workspace, tmp_path):
    root, registry = workspace
    injection = '<script>alert("document")</script>'
    _, evidence_id = intake(workspace, tmp_path, injection)
    path = save_rows(tmp_path / "records.json", [reviewed_row(evidence_id,
                     sample_name='=HYPERLINK("https://invalid.example","sample")', experiment_name='<img src=x onerror=alert(1)>')])
    assert records_command(registry, path, reviewed=True, out=root / "review") == 0
    page = (root / "review" / "index.html").read_text(encoding="utf-8")
    assert injection not in page and "<img src=x" not in page
    assert "&lt;script&gt;" in page and "&lt;img src=x" in page
    with (root / "review" / "records.csv").open(encoding="utf-8-sig", newline="") as stream:
        row = list(csv.DictReader(stream))[0]
    assert row["sample_name"].startswith("'=HYPERLINK")
    exported = json.loads((root / "review" / "knowledge_snapshot.json").read_text(encoding="utf-8"))
    assert exported["records"][0]["sample_name"].startswith("=HYPERLINK")
    assert exported["evidence"][0]["text"] == injection


def test_same_filename_with_different_content_is_not_merged(workspace, tmp_path):
    root, registry = workspace
    first = make_pptx(tmp_path / "one" / "results.pptx", "SET 001 Sample A")
    second = make_pptx(tmp_path / "two" / "results.pptx", "SET 001 Sample B")
    assert invoke("knowledge-intake", "--registry", registry, "--documents", first, second, "--out", root / "review") == 0
    saved = snapshot(registry)
    assert len(saved["documents"]) == 2
    assert len({row["document_id"] for row in saved["documents"]}) == 2
    assert {row["name"] for row in saved["documents"]} == {"results.pptx"}
    assert {row["text"] for row in saved["evidence"]} == {"SET 001 Sample A", "SET 001 Sample B"}
    assert saved["records"] == []


def test_failed_intake_reports_failure_without_creating_usable_evidence(workspace, tmp_path):
    root, registry = workspace
    source = tmp_path / "broken.pptx"
    source.write_bytes(b"not a ZIP archive")
    assert invoke("knowledge-intake", "--registry", registry, "--documents", source, "--out", root / "review") == 2
    saved = snapshot(registry)
    assert saved["evidence"] == [] and saved["records"] == []
    assert saved["documents"][0]["status"] == "failed"
    assert source.read_bytes() == b"not a ZIP archive"


def test_missing_source_and_registry_are_controlled_errors(workspace, tmp_path):
    root, registry = workspace
    assert invoke("knowledge-intake", "--registry", registry, "--documents", tmp_path / "missing.pdf", "--out", root / "review") == 2
    assert snapshot(registry)["documents"] == []
    missing_registry = tmp_path / "does-not-exist" / "catalogue.sqlite3"
    assert invoke("knowledge-export", "--registry", missing_registry, "--out", tmp_path / "export") == 2
    assert not missing_registry.parent.exists()


def test_unavailable_document_does_not_hide_issues_or_skip_later_valid_input(workspace, tmp_path):
    root, registry = workspace
    source = make_pptx(tmp_path / "available.pptx", "SET 004 Sample available")
    missing = tmp_path / "missing.pdf"
    assert invoke("knowledge-intake", "--registry", registry, "--documents", missing, source,
                  "--out", root / "review") == 2
    saved = snapshot(registry)
    assert len(saved["documents"]) == 1 and saved["documents"][0]["name"] == "available.pptx"
    assert saved["evidence"][0]["text"] == "SET 004 Sample available"
    intake_result = json.loads((root / "review" / "intake_result.json").read_text(encoding="utf-8"))
    results = intake_result["results"] if isinstance(intake_result, dict) else intake_result
    assert len(results) == 2
    failed = next(row for row in results if row["status"] == "failed")
    assert Path(failed["source_path"]) == missing
    assert "source_unavailable" in {issue["code"] for issue in failed["issues"]}
    page = (root / "review" / "index.html").read_text(encoding="utf-8")
    assert "source_unavailable" in page and "missing.pdf" in page
    exported = json.loads((root / "review" / "knowledge_snapshot.json").read_text(encoding="utf-8"))
    assert exported["intake_results"] == results


def test_non_sqlite_registry_is_controlled_error_and_unchanged(tmp_path, capsys):
    path = tmp_path / "not-a-registry.sqlite3"
    path.write_bytes(b"ordinary user document")
    assert invoke("knowledge-export", "--registry", path, "--out", tmp_path / "export") == 2
    assert "ERROR:" in capsys.readouterr().err
    assert path.read_bytes() == b"ordinary user document"


def test_registry_and_generated_review_files_cannot_be_intake_sources(workspace):
    root, registry = workspace
    before = snapshot(registry)
    for source in (registry, root / "index.html", root / "knowledge_snapshot.json"):
        content = source.read_bytes()
        assert invoke("knowledge-intake", "--registry", registry, "--documents", source, "--out", root) == 2
        assert source.read_bytes() == content
    assert snapshot(registry) == before


def test_invalid_later_record_rolls_back_entire_cli_batch(workspace, tmp_path):
    _, registry = workspace
    _, evidence_id = intake(workspace, tmp_path)
    before = snapshot(registry)
    path = save_rows(tmp_path / "records.json", [reviewed_row(evidence_id),
                     reviewed_row("missing-evidence", record_id="rec_B", sample_id="sample_B")])
    assert records_command(registry, path, reviewed=True) == 2
    assert snapshot(registry) == before


def test_export_reads_without_rewriting_registry(workspace, tmp_path):
    _, registry = workspace
    before, modified = registry.read_bytes(), registry.stat().st_mtime_ns
    assert invoke("knowledge-export", "--registry", registry, "--out", tmp_path / "export") == 0
    assert registry.read_bytes() == before
    assert registry.stat().st_mtime_ns == modified
