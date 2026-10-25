import pandas as pd
from pptx import Presentation

from battery_insight import slide_link as sl


def _deck(path, slides):
    prs = Presentation()
    for title, body in slides:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        tb = s.shapes.add_textbox(0, 0, 100, 100)
        tb.text_frame.text = body
    prs.save(str(path))


def test_slide_link(tmp_path):
    root = tmp_path / "slides" / "홍길동"
    root.mkdir(parents=True)
    _deck(root / "260720 SJ1300B 평가.pptx", [
        ("SJ1300B MT-1 cycle 결과", "260707 Set1 SJ1300B MT-1 45도 cycle 용량 retention 90% mAh"),
        ("개요", "목차"),
        ("신규 소재 XZ77Q 용량", "XZ77Q 사이클 retention 효율 mAh 결과"),
    ])
    (root / "~$tmp.pptx").write_bytes(b"x")
    (root / "broken.pptx").write_bytes(b"not a zip")

    proj = tmp_path / "proj"
    (proj / "dx").mkdir(parents=True)
    pd.DataFrame([
        dict(test_key="k1", owner_name="홍길동", test_date="2026-07-07", test_type="Cycle", sample_name="SJ1300B MT-1 45도",
             tags="material=SJ1300B; material=MT-1", test_name="07005399_260707_Set1_SJ1300B_MT-1", context_path="홍길동",
             temperature_c=45.0, n_channels=4),
        dict(test_key="k2", owner_name="홍길동", test_date="2026-03-01", test_type="RPT", sample_name="QQ9ZZ 샘플",
             tags="material=QQ9ZZ", test_name="07005399_260301_Set1_QQ9ZZ", context_path="홍길동", n_channels=2),
    ]).to_csv(proj / "dx" / "dx_tests.csv", index=False)

    idx = sl.index_slides(tmp_path / "slides", tmp_path / "c.parquet", workers=2, log=lambda *_: None)
    assert idx["file_path"].nunique() == 2 and (idx["error"] != "").sum() == 1
    assert (idx["slide_no"] > 0).sum() == 3
    again = sl.index_slides(tmp_path / "slides", tmp_path / "c.parquet", workers=2, log=lambda *_: None)
    assert len(again) == len(idx)

    pairs = sl.match_tests(proj, idx)
    assert set(pairs["test_key"]) == {"k1"}
    assert pairs.iloc[0]["slide_no"] == 1 and pairs.iloc[0]["confidence"] == "matched", pairs.iloc[0].to_dict()

    res = sl.build_slide_link(proj, tmp_path / "slides", workers=2, log=lambda *_: None)
    assert list(res["matched"]["test_key"]) == ["k1"]
    assert list(res["tests_without_slides"]["test_key"]) == ["k2"]
    sa = res["slides_without_data"]
    assert list(sa["page"]) == [3], (res["links"].to_dict("records"), idx[["slide_no","n_result_words"]].values.tolist())
    assert (proj / "slides" / "slide_link_report.xlsx").exists()
