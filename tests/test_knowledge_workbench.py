"""Slide evidence -> reviewed records -> raw links -> unchanged numeric data."""
import copy
import json
from types import SimpleNamespace

import pandas as pd
import pytest

from battery_insight.database_workbench import build_database_bundle
from battery_insight.document_intake import extract_document
from battery_insight.knowledge_registry import KnowledgeRegistry
from battery_insight.raw_linking import propose_links, review_links
from test_database_workbench import database_fixture, write_json


def knowledge_fixture(tmp_path, *, physical_alias=False, composition_basis=None,
                      composition=False):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    config, database, output, raw_rows = database_fixture(tmp_path)
    config_data = json.loads(config.read_text())
    config_data["design"]["fields"] = [
        {"field_id": "cathode_active_mass_loading_mg_cm2", "label": "Cathode active loading",
         "kind": "numeric", "unit": "mg/cm2", "scale": 10, "group": "cathode_loading"},
    ]
    components = ["cathode_active_wt_pct", "cathode_conductive_wt_pct", "cathode_binder_wt_pct"]
    if composition:
        config_data["design"]["fields"].append({
            "field_id": "cathode_mix", "label": "Cathode composition", "kind": "composition",
            "unit": "wt%", "composition_basis": "electrode_total", "components": components,
            "group": "cathode_composition"})
    write_json(config, config_data)
    baseline = build_database_bundle(config, output)

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    title = slide.shapes.add_textbox(0, 0, Inches(8), Inches(1))
    title.text = "합성 검증 자료 — 실제 실험 아님\n시험 시작일 2026-10-22 / SET 001 / RPT BOL"
    table = slide.shapes.add_table(3, 5, 0, Inches(1), Inches(9), Inches(2)).table
    for column, value in enumerate(["Sample", "Design", "Batch", "Loading mg/cm2", "Physical cell"]):
        table.cell(0, column).text = value
    for index in range(2):
        for column, value in enumerate([f"합성 시료 {index + 1}", "D1", "B1", str(8.5 + index),
                                        "physical-001" if physical_alias else "미확인"]):
            table.cell(index + 1, column).text = value
    source = tmp_path / "synthetic_slides.pptx"
    presentation.save(source)
    extraction = extract_document(source)
    assert extraction["status"] == "complete"

    registry_path = tmp_path / "knowledge.sqlite"
    records = []
    with KnowledgeRegistry(registry_path) as registry:
        registry.import_document(extraction)
        assert registry.list_records() == []  # Import alone cannot assert a sample identity.
        for index in range(2):
            evidence_ids = [block["block_id"] for block in extraction["blocks"]
                            if block["locator"] == "slide:1/shape:1"
                            or f"/table:r{index + 2}:" in block["locator"]]
            metadata = {"temperature_c": 25, "pressure_mpa": 2.5,
                        "cathode_active_mass_loading_mg_cm2": 8.5 + index}
            if composition:
                metadata.update(dict(zip(components, (96, 2, 2))))
                if composition_basis:
                    metadata["cathode_composition_basis"] = composition_basis
            record = {"record_id": f"record-{index}", "status": "confirmed",
                      "sample_id": f"sample-{index}", "sample_name": f"합성 시료 {index + 1}",
                      "experiment_name": "합성 RPT 검증", "test_date": "2026-10-22",
                      "date_kind": "test_start", "namespace": "synthetic-test",
                      "set_no": "001", "test_type": "RPT", "rpt_stage": "BOL",
                      "design_id": "D1", "batch_id": "B1", "metadata": metadata,
                      "evidence_ids": evidence_ids,
                      "field_evidence": {"sample_name": evidence_ids}}
            if physical_alias:
                record["physical_cell_id"] = "physical-001"
            records.extend(registry.upsert_records([record], reviewer="synthetic reviewer",
                                                   reason="Checked the synthetic slide table"))
        keys = [{"experiment_id": raw["experiment_id"], "test_date": "2026-10-22",
                 "date_kind": "test_start", "namespace": "synthetic-test", "set_no": "001",
                 "sample_name": record["sample_name"]}
                for raw, record in zip(raw_rows, records)]
        proposals = propose_links(database, registry, keys)
        assert {row["status"] for row in proposals["rows"]} == {"unique_candidate"}
        decisions = [{"experiment_id": raw["experiment_id"], "record_id": record["record_id"],
                      "expected_record_revision": record["revision"], "source_sha256": raw["source_sha256"],
                      "decision": "confirm"} for raw, record in zip(raw_rows, records)]
        overlay = review_links(database, registry, proposals, decisions,
                               "synthetic reviewer", "Verified date, set and sample for each raw channel")
    overlay_path = tmp_path / "reviewed_links.json"
    write_json(overlay_path, overlay)
    config_data["knowledge_links"] = str(overlay_path)
    write_json(config, config_data)
    return SimpleNamespace(config=config, database=database, output=output, rows=raw_rows,
                           baseline=baseline, source=source, extraction=extraction,
                           registry_path=registry_path, records=records, keys=keys,
                           proposals=proposals, decisions=decisions, overlay=overlay,
                           overlay_path=overlay_path, components=components)


def test_slide_to_database_link_keeps_numeric_data_and_source_provenance(tmp_path):
    fixture = knowledge_fixture(tmp_path)
    immutable = {path: path.read_bytes() for path in fixture.database.rglob("*") if path.is_file()}
    source_before = fixture.source.read_bytes()
    bundle = build_database_bundle(fixture.config, fixture.output)
    assert bundle["observations"] == fixture.baseline["observations"]
    assert bundle["knowledge_summary"]["linked_experiments"] == 2
    assert bundle["knowledge_summary"]["issues"] == []
    for index, (row, raw, record) in enumerate(zip(bundle["experiments"], fixture.rows, fixture.records)):
        assert row["label"] == record["sample_name"]
        assert row["cell_id"] == row["source_cell_id"] == raw["cell_id"]
        assert row["cell_identity_status"] == "explicit"
        for field in ("source_id", "source_sha256", "test_id", "run_path", "raw_path"):
            assert row[field] == raw[field]
        assert row["cathode_active_mass_loading_mg_cm2"] == 8.5 + index
        assert isinstance(row["cathode_active_mass_loading_mg_cm2"], (int, float))
        assert row["pressure_mpa"] == 2.5
        link = row["knowledge_link"]
        assert link["set_no"] == "001"
        assert link["evidence_ids"] == record["evidence_ids"]
        assert {item["document_id"] for item in link["evidence_locations"]} == {fixture.extraction["document_id"]}
        assert all(item["locator"].startswith("slide:1/") for item in link["evidence_locations"])
    assert all(path.read_bytes() == encoded for path, encoded in immutable.items())
    assert fixture.source.read_bytes() == source_before


def test_physical_alias_only_applies_after_explicit_raw_link_review(tmp_path):
    fixture = knowledge_fixture(tmp_path, physical_alias=True)
    pending = copy.deepcopy(fixture.overlay)
    pending["links"] = []
    write_json(fixture.overlay_path, pending)
    unlinked = build_database_bundle(fixture.config, fixture.output)
    assert {row["cell_id"] for row in unlinked["experiments"]} == {"c0", "c1"}
    assert unlinked["knowledge_summary"]["linked_experiments"] == 0
    write_json(fixture.overlay_path, fixture.overlay)
    linked = build_database_bundle(fixture.config, fixture.output)
    assert {row["cell_id"] for row in linked["experiments"]} == {"physical-001"}
    assert {row["source_cell_id"] for row in linked["experiments"]} == {"c0", "c1"}
    assert {row["cell_identity_status"] for row in linked["experiments"]} == {"mapped_physical_alias"}
    assert linked["observations"] == fixture.baseline["observations"]
    for raw in fixture.rows:
        gold = pd.read_parquet(raw["run_path"] + "/gold/cycles.parquet")
        assert set(gold.cell_id) == {raw["cell_id"]}


def test_unrelated_record_append_keeps_existing_links_but_revised_record_is_skipped(tmp_path):
    fixture = knowledge_fixture(tmp_path)
    with KnowledgeRegistry(fixture.registry_path) as registry:
        new_record = {key: value for key, value in fixture.records[0].items()
                      if key not in {"revision", "content_hash", "created_at", "updated_at"}}
        new_record.update(record_id="unrelated", sample_id="sample-unrelated",
                          sample_name="합성 다른 시료", set_no="002")
        registry.upsert_records([new_record], reviewer="reviewer", reason="Independent later slide")
    unchanged = build_database_bundle(fixture.config, fixture.output)
    assert unchanged["knowledge_summary"]["linked_experiments"] == 2
    with KnowledgeRegistry(fixture.registry_path) as registry:
        registry.upsert_records([{"record_id": fixture.records[0]["record_id"], "notes": "Revised evidence interpretation"}],
                                reviewer="reviewer", reason="Document correction")
    stale = build_database_bundle(fixture.config, fixture.output)
    assert stale["knowledge_summary"]["linked_experiments"] == 1
    assert stale["experiments"][0]["label"] == fixture.rows[0]["label"]
    assert stale["experiments"][1]["label"] == fixture.records[1]["sample_name"]
    assert stale["observations"] == fixture.baseline["observations"]
    assert any(item["reason"] == "knowledge_link_stale_record" and item["experiment_id"] == "e0"
               for item in stale["knowledge_summary"]["issues"])


@pytest.mark.parametrize("basis, applied", [(None, False), ("active_material_total", False),
                                           ("electrode_total", True)])
def test_composition_only_compares_matching_explicit_mass_basis(tmp_path, basis, applied):
    fixture = knowledge_fixture(tmp_path, composition=True, composition_basis=basis)
    bundle = build_database_bundle(fixture.config, fixture.output)
    for row in bundle["experiments"]:
        if applied:
            assert [row[key] for key in fixture.components] == [96, 2, 2]
            assert not row["knowledge_link"].get("unapplied_metadata_fields")
        else:
            assert all(row.get(key) is None for key in fixture.components)
            assert row["knowledge_link"]["unapplied_metadata_fields"] == sorted(fixture.components)
    basis_issues = [issue for issue in bundle["knowledge_summary"]["issues"]
                    if issue["reason"] == "knowledge_composition_basis_unverified"]
    assert len(basis_issues) == (0 if applied else 2)
    with KnowledgeRegistry(fixture.registry_path, read_only=True) as registry:
        assert registry.get_record("record-0")["metadata"]["cathode_active_wt_pct"] == 96
    assert bundle["observations"] == fixture.baseline["observations"]


def test_existing_csv_metadata_conflict_does_not_silently_overwrite_user_values(tmp_path):
    fixture = knowledge_fixture(tmp_path)
    path = fixture.database / "experiments.csv"
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    frame.loc[0, "cathode_active_mass_loading_mg_cm2"] = "42"
    frame.to_csv(path, index=False, encoding="utf-8-sig")
    encoded = path.read_bytes()
    bundle = build_database_bundle(fixture.config, fixture.output)
    assert bundle["knowledge_summary"]["linked_experiments"] == 1
    assert bundle["experiments"][0]["cathode_active_mass_loading_mg_cm2"] == 42.0
    assert bundle["experiments"][0]["label"] == fixture.rows[0]["label"]
    assert any(issue["reason"] == "knowledge_link_existing_metadata_conflict"
               for issue in bundle["knowledge_summary"]["issues"])
    assert path.read_bytes() == encoded


def test_record_revision_changing_during_bundle_build_refuses_mixed_snapshot(tmp_path, monkeypatch):
    import battery_insight.database_workbench as adapter

    fixture = knowledge_fixture(tmp_path)
    original = adapter.mapped_cycles
    changed = False

    def change_record_during_read(*args, **kwargs):
        nonlocal changed
        if not changed:
            changed = True
            with KnowledgeRegistry(fixture.registry_path) as registry:
                registry.upsert_records([{"record_id": "record-0", "notes": "Concurrent revision"}],
                                        reviewer="reviewer", reason="Concurrent document review")
        return original(*args, **kwargs)

    monkeypatch.setattr(adapter, "mapped_cycles", change_record_during_read)
    with pytest.raises(ValueError, match="Document knowledge or reviewed links changed"):
        build_database_bundle(fixture.config, fixture.output)
