from dataclasses import replace
from datetime import date
import hashlib
from io import BytesIO
import json
import zipfile

import pytest

from battery_insight.document_intake import IntakeLimits, extract_document


def issue_codes(result):
    return {issue["code"] for issue in result["issues"]}


def write_image(path):
    Image = pytest.importorskip("PIL.Image")
    Image.new("RGB", (12, 10), "white").save(path)
    return path


def write_pptx(path, image_path=None):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, Inches(3), Inches(1)).text = "2026-10-22 SET 03\nSample A"
    table = slide.shapes.add_table(2, 2, 0, Inches(1), Inches(4), Inches(2)).table
    table.cell(0, 0).text = "Set"
    table.cell(0, 1).text = "Sample"
    table.cell(1, 0).text = "03"
    table.cell(1, 1).text = "A"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(1), Inches(1)).text = "Nested label"
    slide.notes_slide.notes_text_frame.text = "Operator note: date is test date, not slide date"
    if image_path is not None:
        slide.shapes.add_picture(str(image_path), 0, Inches(3), width=Inches(1))
    presentation.save(path)
    return path


def write_pdf(path, *, blank=False, encrypted=False):
    pypdf = pytest.importorskip("pypdf")
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    if not blank:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 10 100 Td (SET 03 Sample A) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("not-a-real-password")
    with path.open("wb") as handle:
        writer.write(handle)
    return path


def test_pptx_native_text_table_nested_shape_notes_and_stable_ids(tmp_path):
    path = write_pptx(tmp_path / "report.pptx")
    result = extract_document(path)
    assert result["status"] == "complete"
    assert result["extraction_complete"] is True
    assert result["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert result["document_id"] == "doc_" + result["sha256"][:24]
    blocks = {block["locator"]: block for block in result["blocks"]}
    assert blocks["slide:1/shape:1"]["text"].startswith("2026-10-22")
    assert blocks["slide:1/shape:2/table:r2:c2"]["text"] == "A"
    assert blocks["slide:1/shape:3/shape:1"]["text"] == "Nested label"
    assert blocks["slide:1/notes"]["kind"] == "speaker_notes"
    assert all(block["needs_review"] for block in blocks.values())
    assert len({block["block_id"] for block in blocks.values()}) == len(blocks)
    assert extract_document(path) == result
    copy_path = tmp_path / "renamed.pptx"
    copy_path.write_bytes(path.read_bytes())
    copy = extract_document(copy_path)
    assert copy["document_id"] == result["document_id"]
    assert copy["blocks"] == result["blocks"]
    json.dumps(result, allow_nan=False)


def test_image_and_pptx_image_are_explicit_ocr_pending(tmp_path):
    path = write_image(tmp_path / "scan.png")
    image_result = extract_document(path)
    assert image_result["status"] == "partial"
    assert image_result["metadata"]["width"] == 12
    assert image_result["blocks"][0]["text"] == ""
    assert "ocr_required" in issue_codes(image_result)
    result = extract_document(write_pptx(tmp_path / "report.pptx", path))
    assert result["status"] == "partial"
    assert "ocr_required" in issue_codes(result)
    assert any(block["kind"] == "text" and block["text"] for block in result["blocks"])


def test_xlsx_formulas_dates_and_merged_cells_are_not_inferred(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Experiments"
    sheet["A1"] = "SET 03"
    sheet.merge_cells("A1:B1")
    sheet["C1"] = "=SUM(1,2)"
    sheet["D1"] = date(2026, 10, 22)
    sheet["A2"] = "Sample A"
    sheet["B2"] = 1.25
    path = tmp_path / "registry.xlsx"
    workbook.save(path)
    result = extract_document(path)
    assert result["status"] == "complete"
    blocks = {block["locator"]: block for block in result["blocks"]}
    assert "sheet:Experiments/cell:B1" not in blocks
    assert blocks["sheet:Experiments/cell:C1"]["text"] == "=SUM(1,2)"
    assert blocks["sheet:Experiments/cell:C1"]["kind"] == "formula"
    assert blocks["sheet:Experiments/cell:D1"]["text"].startswith("2026-10-22")
    assert blocks["sheet:Experiments/cell:B2"]["text"] == "1.25"
    assert "formula_not_evaluated" in issue_codes(result)


def test_pdf_text_and_blank_page_have_review_limitations(tmp_path):
    result = extract_document(write_pdf(tmp_path / "report.pdf"))
    assert result["status"] == "partial"
    assert result["blocks"][0]["locator"] == "page:1"
    assert "SET 03 Sample A" in result["blocks"][0]["text"]
    assert "pdf_visual_content_unverified" in issue_codes(result)
    blank = extract_document(write_pdf(tmp_path / "scan.pdf", blank=True))
    assert blank["status"] == "partial"
    assert blank["blocks"][0]["text"] == ""
    assert "ocr_required" in issue_codes(blank)


def test_encrypted_pdf_is_clear_failure(tmp_path):
    result = extract_document(write_pdf(tmp_path / "encrypted.pdf", encrypted=True))
    assert result["status"] == "failed"
    assert result["blocks"] == []
    assert "encrypted_pdf" in issue_codes(result)


def test_pdf_decoded_content_budget_discards_partial_text(tmp_path):
    path = write_pdf(tmp_path / "report.pdf")
    result = extract_document(path, limits=replace(IntakeLimits(), max_pdf_decoded_bytes=10))
    assert result["status"] == "failed"
    assert result["blocks"] == []
    assert "pdf_decoded_limit" in issue_codes(result)


def test_slide_page_limit_and_empty_document(tmp_path):
    pptx = pytest.importorskip("pptx")
    presentation = pptx.Presentation()
    empty = tmp_path / "empty.pptx"
    presentation.save(empty)
    result = extract_document(empty)
    assert result["status"] == "partial"
    assert result["blocks"][0]["kind"] == "empty_document"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = tmp_path / "two.pptx"
    presentation.save(path)
    result = extract_document(path, limits=replace(IntakeLimits(), max_pages=1))
    assert result["status"] == "failed"
    assert "page_limit" in issue_codes(result)


def test_chart_values_are_not_inferred_and_external_hyperlinks_not_fetched(tmp_path, monkeypatch):
    pptx = pytest.importorskip("pptx")
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    import socket

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    chart_data = CategoryChartData()
    chart_data.categories = ["SET 03"]
    chart_data.add_series("capacity", [3.14159])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(3), Inches(2), chart_data)
    run = slide.shapes.add_textbox(0, Inches(3), Inches(3), Inches(1)).text_frame.paragraphs[0].add_run()
    run.text = "External reference"
    run.hyperlink.address = "https://example.invalid/private-report"
    path = tmp_path / "chart.pptx"
    presentation.save(path)

    def no_network(*args, **kwargs):
        raise AssertionError("document intake must not open network connections")

    monkeypatch.setattr(socket, "create_connection", no_network)
    result = extract_document(path)
    assert result["status"] == "partial"
    assert "chart_review_required" in issue_codes(result)
    assert "external_reference_ignored" in issue_codes(result)
    assert "3.14159" not in " ".join(block["text"] for block in result["blocks"])


@pytest.mark.parametrize("suffix", ["pptx", "pdf", "xlsx", "png"])
def test_malformed_documents_fail_without_partial_evidence(tmp_path, suffix):
    path = tmp_path / ("bad." + suffix)
    path.write_bytes(b"this is not a document")
    result = extract_document(path)
    assert result["status"] == "failed"
    assert not result["extraction_complete"]
    assert not result["blocks"]
    assert result["issues"]


@pytest.mark.parametrize("changes, expected", [
    ({"max_file_bytes": 10}, "file_limit"),
    ({"max_text_chars": 5}, "text_limit"),
    ({"max_blocks": 2}, "block_limit"),
    ({"max_archive_uncompressed_bytes": 50}, "archive_expansion_limit"),
    ({"max_archive_members": 2}, "archive_member_limit"),
])
def test_size_text_block_and_archive_limits(tmp_path, changes, expected):
    path = write_pptx(tmp_path / "report.pptx")
    result = extract_document(path, limits=replace(IntakeLimits(), **changes))
    assert result["status"] == "failed"
    assert expected in issue_codes(result)
    assert result["blocks"] == []


def test_sheet_dimension_limit_prevents_empty_cell_expansion(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    workbook.active["XFD1048576"] = "far away"
    path = tmp_path / "sparse.xlsx"
    workbook.save(path)
    result = extract_document(path)
    assert result["status"] == "failed"
    assert "sheet_cell_limit" in issue_codes(result)


def test_archive_dtd_and_unsafe_path_are_rejected(tmp_path):
    for label, name, text, code in [
        ("entity", "sheet.xml", b'<!DOCTYPE sheet [<!ENTITY x SYSTEM "file:///secret">]><sheet/>', "xml_entities_disallowed"),
        ("path", "../sheet.xml", b"<sheet/>", "unsafe_archive"),
    ]:
        path = tmp_path / f"{label}.xlsx"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(name, text)
        result = extract_document(path)
        assert result["status"] == "failed"
        assert code in issue_codes(result)


def test_modification_during_read_discards_all_evidence(tmp_path, monkeypatch):
    import battery_insight.document_intake as module

    path = write_pptx(tmp_path / "report.pptx")
    original = module._extract_pptx

    def extracting_with_change(data, out):
        original(data, out)
        path.write_bytes(path.read_bytes() + b"changed")

    monkeypatch.setattr(module, "_extract_pptx", extracting_with_change)
    result = extract_document(path)
    assert result["status"] == "failed"
    assert result["blocks"] == []
    assert "source_changed" in issue_codes(result)


def test_missing_dependency_and_source_are_structured(tmp_path, monkeypatch):
    import battery_insight.document_intake as module

    path = write_pptx(tmp_path / "report.pptx")

    def missing(data, out):
        raise ModuleNotFoundError("pptx")

    monkeypatch.setattr(module, "_extract_pptx", missing)
    assert "missing_dependency" in issue_codes(extract_document(path))
    assert "source_unavailable" in issue_codes(extract_document(tmp_path / "missing.pdf"))
    assert "local_files_only" in issue_codes(extract_document("https://example.invalid/report.pdf"))
    unsupported = tmp_path / "active.xlsm"
    unsupported.write_bytes(b"test")
    assert "unsupported_format" in issue_codes(extract_document(unsupported))


@pytest.mark.parametrize("value", [0, -1, 1.5, True])
def test_limits_must_be_positive_integers(value):
    with pytest.raises(ValueError):
        IntakeLimits(max_pages=value)
