"""Synthetic PPT -> reviewed manufacturing -> cell/tests -> CYC/Parquet demo."""
from pathlib import Path
import json


def create_demo(destination, *, reviewed=True):
    from pptx import Presentation
    from pptx.util import Inches
    from . import library
    from . import traceability as trace
    from .native_cyc.selftest import fixture_4101
    from .slide_candidates import extract_proposals, review_proposal
    root = Path(destination).resolve()
    if root.exists() and any(root.iterdir()):
        raise ValueError("합성 예제는 비어 있는 새 프로젝트에만 만듭니다.")
    root.mkdir(parents=True, exist_ok=True)
    slides = Presentation()
    slides.slide_width, slides.slide_height = Inches(16), Inches(9)
    def table_slide(title, headers, rows):
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        slide.shapes.add_textbox(Inches(.3), Inches(.2), Inches(15), Inches(.8)).text = "합성 예시 · 실제 실험 자료 아님 · " + title
        table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(.3), Inches(1.3), Inches(15), Inches(5.8)).table
        for r, row in enumerate([headers, *rows]):
            for c, value in enumerate(row):
                table.cell(r, c).text = str(value)
    headers = ["항목종류", "항목명", "소재로트ID", "전극로트ID", "물리셀ID", "공정ID", "공정종류", "공정일", "조성기준", "활물질비율(%)", "로딩(mg/cm²)", "압착압(MPa)"]
    table_slide("소재·공정·셀 정보", headers, [
        ["material_lot", "합성 활물질 로트", "SYN-AM-001", "", "", "", "", "2026-01-01", "", "", "", ""],
        ["process_run", "합성 배합", "", "", "", "SYN-MIX-001", "배합", "2026-01-02", "전극 전체", "65", "", ""],
        ["electrode_lot", "합성 양극 로트", "", "SYN-EL-001", "", "", "", "2026-01-02", "단면", "", "6.3", ""],
        ["process_run", "합성 조립", "", "", "", "SYN-ASM-001", "조립", "2026-01-03", "", "", "", ""],
        ["cell", "합성 물리 셀", "", "", "SYN-CELL-001", "", "", "2026-01-03", "", "", "", ""],
        ["process_run", "합성 WIP", "", "", "", "SYN-WIP-001", "WIP", "2026-01-10", "", "", "", "100"],
    ])
    table_slide("동일 셀의 WIP 전·후 시험", ["샘플ID", "샘플명", "물리셀ID", "시험명", "시험시작일", "세트번호", "프로젝트", "시험종류", "RPT시점", "설계ID", "배치ID", "온도(°C)", "구동압(MPa)"], [
        ["SYN-SAMPLE-001", "합성 샘플 A", "SYN-CELL-001", "합성 BOL", "2026-01-05", "001", "synthetic_demo", "RPT", "BOL", "SYN-DESIGN", "SYN-BATCH", "25", "1"],
        ["SYN-SAMPLE-001", "합성 샘플 A", "SYN-CELL-001", "합성 WIP 이후", "2026-01-11", "001", "synthetic_demo", "RPT", "after_WIP", "SYN-DESIGN", "SYN-BATCH", "25", "1"],
    ])
    inputs = root / "inputs"
    inputs.mkdir()
    ppt = inputs / "합성_공정_시험_연결.pptx"
    slides.save(ppt)
    library.import_documents(root, [ppt])
    proposals = extract_proposals(root)
    datasets = {}
    for stage in ("BOL", "after_WIP"):
        raw = root / "inputs" / stage
        fixture_4101(raw)
        result = library.convert_and_register(root, raw)
        datasets[stage] = result["output"]
    if reviewed:
        snap = trace.snapshot(root)
        for candidate in snap["proposals"]:
            payload = {**candidate["payload"], "status": "confirmed"}
            review_proposal(root, candidate["proposal_id"], payload, reviewer="synthetic_demo", reason="합성 예제의 명시적 대응표 확인")
        snap = trace.snapshot(root)
        entities = snap["entities"]
        by_key = {}
        for entity in entities:
            attrs = entity["attributes"]
            key = attrs.get("process_run_id") or attrs.get("physical_cell_id") or attrs.get("lot_id")
            by_key[key] = entity
        records = {row["rpt_stage"]: row for row in snap["records"]}
        evidence = lambda *keys: sorted({ev for key in keys for ev in by_key[key]["evidence_ids"]})
        definitions = [
            ("SYN-AM-001", "SYN-MIX-001", "consumed_in", "양극 활물질"),
            ("SYN-MIX-001", "SYN-EL-001", "produces", "양극 전극"),
            ("SYN-EL-001", "SYN-ASM-001", "consumed_in", "양극 전극"),
            ("SYN-ASM-001", "SYN-CELL-001", "produces", "조립 셀"),
            ("SYN-CELL-001", "SYN-WIP-001", "processed_in", "동일 셀 후처리"),
        ]
        edges = [{"edge_id": "syn_edge_" + str(index), "source_id": by_key[left]["entity_id"],
                  "target_id": by_key[right]["entity_id"], "relation": relation, "role": role,
                  "status": "confirmed", "evidence_ids": evidence(left, right)}
                 for index, (left, right, relation, role) in enumerate(definitions)]
        for stage, record in records.items():
            edges.append({"edge_id": "syn_test_" + stage, "source_id": by_key["SYN-CELL-001"]["entity_id"],
                          "target_id": record["record_id"], "target_kind": "record", "relation": "tested_as", "status": "confirmed",
                          "evidence_ids": record["evidence_ids"]})
        edges.append({"edge_id": "syn_wip_after", "source_id": by_key["SYN-WIP-001"]["entity_id"],
                      "target_id": records["after_WIP"]["record_id"], "target_kind": "record", "relation": "conditioned_for",
                      "status": "confirmed", "evidence_ids": records["after_WIP"]["evidence_ids"]})
        trace.save_bundle(root, {"edges": edges}, reviewer="synthetic_demo", reason="합성 공정·셀·시험 연결 확인")
        for row in library.library_rows(root):
            stage = next(stage for stage, output in datasets.items() if Path(row["path"]).is_relative_to(Path(output)))
            library.link_record(root, row["dataset_id"], records[stage]["record_id"], reviewer="synthetic_demo", reason="합성 CYC 대응표 확인")
    (root / "SYNTHETIC_DEMO.txt").write_text("합성 동작 예제입니다. 실제 공정·설계·시험 조건을 의미하지 않습니다.\n일부 공정과 투입 소재는 개념 설명을 위해 생략했습니다.", encoding="utf-8")
    proposals["pending"] = sum(item["state"] == "pending" for item in trace.snapshot(root)["proposals"])
    return {"status": "success", "project": str(root), "ppt": str(ppt), "reviewed_synthetic": reviewed,
            **proposals, "policy": "synthetic_reviewed_demo" if reviewed else proposals["policy"]}
