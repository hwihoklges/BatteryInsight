"""Read local slide/report evidence without inferring experiment identities.

Only native document text is extracted. Images, charts, ambiguous table layout,
and formula results remain explicit review tasks. Nothing is written or fetched.
"""
from __future__ import annotations

from dataclasses import dataclass, fields
from datetime import date, datetime, time
import hashlib
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import Any
import warnings
import zipfile


@dataclass(frozen=True)
class IntakeLimits:
    max_file_bytes: int = 100 * 1024 * 1024
    max_pages: int = 2000
    max_text_chars: int = 10_000_000
    max_blocks: int = 100_000
    max_archive_uncompressed_bytes: int = 256 * 1024 * 1024
    max_archive_members: int = 100_000
    max_sheet_cells: int = 1_000_000
    max_pdf_stream_bytes: int = 32 * 1024 * 1024
    max_pdf_decoded_bytes: int = 256 * 1024 * 1024

    def __post_init__(self):
        for field in fields(self):
            value = getattr(self, field.name)
            if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
                raise ValueError(f"{field.name} must be a positive integer")


class _IntakeError(Exception):
    def __init__(self, code: str, message: str):
        self.code = code
        super().__init__(message)


class _Collector:
    def __init__(self, result: dict, limits: IntakeLimits):
        self.result = result
        self.limits = limits
        self.text_chars = 0
        self.visited = 0

    def visit(self):
        self.visited += 1
        if self.visited > self.limits.max_blocks:
            raise _IntakeError("block_limit", "Document contains too many text/shape elements.")

    def add(self, locator: str, text: str, kind: str):
        if len(self.result["blocks"]) >= self.limits.max_blocks:
            raise _IntakeError("block_limit", "Document exceeds the evidence block limit.")
        self.text_chars += len(text)
        if self.text_chars > self.limits.max_text_chars:
            raise _IntakeError("text_limit", "Document exceeds the extracted text limit.")
        identity = f"{self.result['sha256']}\0{locator}\0{kind}"
        self.result["blocks"].append({
            "block_id": "ev_" + hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24],
            "locator": locator, "text": text, "kind": kind, "needs_review": True,
        })

    def issue(self, code: str, message: str, locator: str | None = None,
              *, incomplete: bool = False):
        issue = {"code": code, "message": message}
        if locator is not None:
            issue["locator"] = locator
        self.result["issues"].append(issue)
        if incomplete:
            self.result["status"] = "partial"
            self.result["extraction_complete"] = False

    def placeholder(self, locator: str, kind: str, code: str, message: str):
        self.add(locator, "", kind)
        self.issue(code, message, locator, incomplete=True)


def _read_bounded(path: Path, limit: int) -> bytes:
    if not path.is_file():
        raise _IntakeError("source_unavailable", "Source must be an existing local file.")
    if path.stat().st_size > limit:
        raise _IntakeError("file_limit", "Source exceeds the allowed file size.")
    with path.open("rb") as handle:
        data = handle.read(limit + 1)
    if len(data) > limit:
        raise _IntakeError("file_limit", "Source exceeds the allowed file size.")
    return data


def _source_digest(path: Path, limit: int) -> str:
    digest = hashlib.sha256()
    size = 0
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            size += len(chunk)
            if size > limit:
                raise _IntakeError("source_changed", "Source grew while it was being read.")
            digest.update(chunk)
    return digest.hexdigest()


def _check_archive(data: bytes, out: _Collector) -> list[str]:
    """Validate every member before any XML parser sees the immutable snapshot."""
    names = []
    expanded = 0
    with zipfile.ZipFile(BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > out.limits.max_archive_members:
            raise _IntakeError("archive_member_limit", "Archive contains too many members.")
        if sum(info.file_size for info in members) > out.limits.max_archive_uncompressed_bytes:
            raise _IntakeError("archive_expansion_limit", "Archive exceeds the uncompressed size limit.")
        seen = set()
        for info in members:
            name = info.filename
            path = PurePosixPath(name)
            if (name in seen or "\\" in name or path.is_absolute() or ".." in path.parts
                    or any(":" in part for part in path.parts)):
                raise _IntakeError("unsafe_archive", "Archive has duplicate or unsafe member paths.")
            seen.add(name)
            names.append(name)
            if info.flag_bits & 1:
                raise _IntakeError("encrypted_archive", "Encrypted Office archives are not supported.")
            tail = b""
            is_xml = name.lower().endswith((".xml", ".rels"))
            with archive.open(info) as handle:
                while chunk := handle.read(1024 * 1024):
                    expanded += len(chunk)
                    if expanded > out.limits.max_archive_uncompressed_bytes:
                        raise _IntakeError("archive_expansion_limit", "Archive expansion exceeds the limit.")
                    if is_xml:
                        probe = (tail + chunk).upper().replace(b"\x00", b"")
                        if b"<!DOCTYPE" in probe or b"<!ENTITY" in probe:
                            raise _IntakeError("xml_entities_disallowed", "XML entity/DTD declarations are not accepted.")
                        if b'TARGETMODE="EXTERNAL"' in probe or b"TARGETMODE='EXTERNAL'" in probe:
                            out.issue("external_reference_ignored", "External document references are not fetched.", name)
                        tail = chunk[-64:]
        if any(name.lower().endswith("vbaproject.bin") for name in names):
            out.issue("active_content_ignored", "Embedded VBA is not executed or extracted.")
    return names


def _extract_pptx(data: bytes, out: _Collector):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    _check_archive(data, out)
    presentation = Presentation(BytesIO(data))
    if len(presentation.slides) > out.limits.max_pages:
        raise _IntakeError("page_limit", "Presentation exceeds the slide limit.")
    out.result["metadata"]["slide_count"] = len(presentation.slides)

    def visit_shapes(shapes, prefix, depth=0):
        if depth > 64:
            raise _IntakeError("nesting_limit", "Presentation shape nesting exceeds 64 levels.")
        for index, shape in enumerate(shapes, 1):
            out.visit()
            locator = f"{prefix}/shape:{index}"
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                visit_shapes(shape.shapes, locator, depth + 1)
            elif shape.has_table:
                for row_index, row in enumerate(shape.table.rows, 1):
                    for col_index, cell in enumerate(row.cells, 1):
                        out.visit()
                        if not cell.is_spanned and cell.text.strip():
                            out.add(f"{locator}/table:r{row_index}:c{col_index}", cell.text, "table_cell")
            elif shape.has_chart:
                out.placeholder(locator, "chart", "chart_review_required",
                                "Chart values are not inferred; review the chart and its source table.")
            elif shape.shape_type in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE):
                out.placeholder(locator, "image", "ocr_required", "Image content needs visual review or OCR.")
            elif shape.has_text_frame and shape.text.strip():
                out.add(locator, shape.text, "text")
            elif not shape.has_text_frame:
                out.placeholder(locator, "embedded_object", "visual_review_required",
                                "Unsupported drawing or embedded object needs visual review.")

    for index, slide in enumerate(presentation.slides, 1):
        prefix = f"slide:{index}"
        count_before = len(out.result["blocks"])
        visit_shapes(slide.shapes, prefix)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                out.add(f"{prefix}/notes", notes.text, "speaker_notes")
        if len(out.result["blocks"]) == count_before:
            out.placeholder(prefix, "empty_slide", "no_native_text",
                            "Slide has no extractable native text; review the original slide.")
    out.issue("layout_not_interpreted", "Slide positions, arrows, colors and table headers do not automatically establish experiment/sample relationships.")


def _extract_pdf(data: bytes, out: _Collector):
    import pypdf

    if not hasattr(pypdf, "apply_configuration"):
        raise _IntakeError("dependency_version", "PDF intake requires pypdf>=6.18 for bounded stream decoding.")
    stream_limit = out.limits.max_pdf_stream_bytes
    with pypdf.apply_configuration(
            maximum_declared_stream_length=stream_limit,
            array_based_stream_maximum_output_length=stream_limit,
            zlib_maximum_output_length=stream_limit,
            lzw_maximum_output_length=stream_limit,
            run_length_maximum_output_length=stream_limit,
            page_tree_maximum_entries=out.limits.max_pages * 4,
            page_tree_maximum_depth=100,
            xform_maximum_invocations_per_extraction=1000,
            jbig2dec_binary=None):
        reader = pypdf.PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise _IntakeError("encrypted_pdf", "Encrypted PDF is not accepted; supply an unlocked local copy.")
        if len(reader.pages) > out.limits.max_pages:
            raise _IntakeError("page_limit", "PDF exceeds the page limit.")
        out.result["metadata"]["page_count"] = len(reader.pages)
        decoded_bytes = 0
        decoded_objects = set()

        def account_stream(stream):
            nonlocal decoded_bytes
            if stream is None:
                return
            stream = stream.get_object()
            key = id(stream)
            if key in decoded_objects:
                return
            decoded_objects.add(key)
            if len(decoded_objects) > out.limits.max_blocks:
                raise _IntakeError("block_limit", "PDF contains too many decoded content streams.")
            if hasattr(stream, "get_data"):
                decoded_bytes += len(stream.get_data())
            else:
                # A page /Contents array may contain multiple streams.
                for member in stream:
                    account_stream(member)
            if decoded_bytes > out.limits.max_pdf_decoded_bytes:
                raise _IntakeError("pdf_decoded_limit", "PDF exceeds the total decoded content limit.")

        visited_forms = set()

        def account_resources(resources, depth=0):
            if depth > 64:
                raise _IntakeError("nesting_limit", "PDF form nesting exceeds 64 levels.")
            resources = resources.get_object() if hasattr(resources, "get_object") else resources
            objects = resources.get("/XObject", {})
            objects = objects.get_object() if hasattr(objects, "get_object") else objects
            for reference in objects.values():
                form = reference.get_object()
                if form.get("/Subtype") != "/Form" or id(form) in visited_forms:
                    continue
                visited_forms.add(id(form))
                account_stream(form)
                account_resources(form.get("/Resources", {}), depth + 1)

        for index, page in enumerate(reader.pages, 1):
            out.visit()
            locator = f"page:{index}"
            account_stream(page.get("/Contents"))
            account_resources(page.get("/Resources", {}))
            text = page.extract_text() or ""
            if text.strip():
                out.add(locator, text, "page_text")
            else:
                out.placeholder(locator, "page_without_text", "ocr_required",
                                "Page has no native text; visual review or OCR is required.")
        out.issue("pdf_layout_review", "PDF text order and table structure require review; diagrams, charts and image text are not extracted.")
        # Even a text-rich PDF can contain scanned legends or plotted values.
        out.issue("pdf_visual_content_unverified", "Text extraction cannot establish that all visual content was captured.", incomplete=True)


def _cell_text(value: Any) -> str:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value)


def _extract_xlsx(data: bytes, out: _Collector):
    from openpyxl import load_workbook

    names = _check_archive(data, out)
    workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        if len(workbook.worksheets) > out.limits.max_pages:
            raise _IntakeError("page_limit", "Workbook exceeds the sheet limit.")
        out.result["metadata"]["sheet_count"] = len(workbook.worksheets)
        examined_cells = 0
        for sheet in workbook.worksheets:
            # Bound the iteration rectangle, including blank cells. Sparse sheets
            # with extreme dimensions are rejected rather than silently truncated.
            rows, cols = sheet.max_row or 1, sheet.max_column or 1
            examined_cells += rows * cols
            if examined_cells > out.limits.max_sheet_cells:
                raise _IntakeError("sheet_cell_limit", "Workbook dimensions exceed the inspected cell limit.")
            prefix = "sheet:" + sheet.title.replace("%", "%25").replace("/", "%2F")
            count_before = len(out.result["blocks"])
            for row in sheet.iter_rows(max_row=rows, max_col=cols):
                for cell in row:
                    if cell.value is None:
                        continue
                    out.visit()
                    locator = f"{prefix}/cell:{cell.coordinate}"
                    formula = cell.data_type == "f"
                    if formula and not isinstance(cell.value, str):
                        formula_text = getattr(cell.value, "text", None)
                        text = formula_text if formula_text is not None else "[unsupported formula representation]"
                    else:
                        text = _cell_text(cell.value)
                    out.add(locator, text, "formula" if formula else "cell")
                    if formula:
                        out.issue("formula_not_evaluated", "Formula text is evidence; cached or calculated results are not used.", locator)
            if len(out.result["blocks"]) == count_before:
                out.placeholder(prefix, "empty_sheet", "no_native_text", "Sheet contains no nonempty cells.")
        for name in names:
            if name.startswith(("xl/charts/", "xl/media/")) and not name.endswith("/"):
                out.placeholder("part:" + name, "embedded_visual", "visual_review_required",
                                "Embedded chart/image is not interpreted; review the source workbook.")
        out.issue("spreadsheet_layout_review", "Merged cells are not forward-filled. Units, headers, hidden rows/sheets and displayed number formatting require review.")
    finally:
        workbook.close()


def _extract_image(data: bytes, out: _Collector):
    from PIL import Image

    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(BytesIO(data)) as picture:
            out.result["metadata"].update({"width": picture.width, "height": picture.height,
                                          "image_format": picture.format})
            picture.verify()
    out.placeholder("image:1", "image", "ocr_required", "Image recorded as evidence; visual review or OCR is required before entering sample information.")


def extract_document(path: str | Path, *, limits: IntakeLimits | None = None) -> dict:
    """Return deterministic native text evidence and explicit review limitations.

    Malformed, unsupported, changed or over-limit documents return ``failed``
    with no evidence blocks. No identifiers, dates or numerical chart values are
    guessed. All extracted blocks need semantic review before registry linking.
    """
    limits = limits or IntakeLimits()
    source = Path(path).expanduser().absolute()
    result = {"schema_version": "document_intake_v1", "document_id": None, "sha256": None,
              "source_path": str(source), "name": source.name,
              "format": source.suffix.lower().lstrip("."), "status": "complete",
              "extraction_complete": True, "blocks": [], "issues": [], "metadata": {}}
    out = _Collector(result, limits)
    try:
        if str(path).lower().startswith(("http://", "https://", "ftp://")):
            raise _IntakeError("local_files_only", "Only local filesystem documents are accepted.")
        data = _read_bounded(source, limits.max_file_bytes)
        result["sha256"] = hashlib.sha256(data).hexdigest()
        result["document_id"] = "doc_" + result["sha256"][:24]
        result["metadata"]["file_bytes"] = len(data)
        kind = result["format"]
        if kind == "pptx":
            _extract_pptx(data, out)
        elif kind == "pdf":
            _extract_pdf(data, out)
        elif kind == "xlsx":
            _extract_xlsx(data, out)
        elif kind in {"png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp"}:
            _extract_image(data, out)
        else:
            raise _IntakeError("unsupported_format", "Supported formats: PPTX, PDF, XLSX, PNG, JPEG, TIFF, BMP and WebP.")
        if not result["blocks"]:
            out.placeholder("document:1", "empty_document", "no_native_text", "Document contains no extractable native text.")
        if _source_digest(source, limits.max_file_bytes) != result["sha256"]:
            raise _IntakeError("source_changed", "Source changed during extraction; retry with a stable copy.")
    except Exception as error:
        result["status"] = "failed"
        result["extraction_complete"] = False
        result["blocks"] = []
        if isinstance(error, _IntakeError):
            code = error.code
            message = str(error)
        elif isinstance(error, ImportError):
            code = "missing_dependency"
            message = "Install battery-insight[documents] to enable this document format. " + str(error)
        else:
            code = "document_read_error"
            message = f"{type(error).__name__}: {error}"
        out.issue(code, message)
    return result
