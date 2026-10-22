"""Local workbench data and bounded, provenance-checked evidence readers."""
from __future__ import annotations

import csv
from datetime import date, datetime, timezone
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import tempfile
import zipfile

import numpy as np
import pandas as pd

from .features import DEFAULTS, _curve_features, _direction, _segments
from .ingest import sha256_file

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_UPLOAD_ROWS = 10000
MAX_UPLOAD_COLUMNS = 200
MAX_TABLE_ROWS = 500000
MAX_CYCLE_POINTS = 200000
MAX_SCAN_ROWS = 2000000
MAX_CURVE_POINTS = 1200
CONTROL_FIELDS = [
    {"field_id": "temperature_c", "label": "Temperature", "kind": "numeric", "unit": "degC"},
    {"field_id": "pressure_mpa", "label": "Pressure", "kind": "numeric", "unit": "MPa"},
    {"field_id": "charge_rate_c", "label": "Charge rate", "kind": "numeric", "unit": "C"},
    {"field_id": "discharge_rate_c", "label": "Discharge rate", "kind": "numeric", "unit": "C"},
    {"field_id": "voltage_min_v", "label": "Lower voltage limit", "kind": "numeric", "unit": "V"},
    {"field_id": "voltage_max_v", "label": "Upper voltage limit", "kind": "numeric", "unit": "V"},
    {"field_id": "formation_id", "label": "Formation protocol", "kind": "categorical", "unit": ""},
    {"field_id": "batch_id", "label": "Manufacturing batch", "kind": "categorical", "unit": ""},
]


def json_safe(value):
    if isinstance(value, dict):
        return {str(k): json_safe(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [json_safe(v) for v in value]
    if isinstance(value, (np.generic,)):
        value = value.item()
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()
    if value is None or (isinstance(value, float) and not math.isfinite(value)):
        return None
    if value is pd.NA or value is pd.NaT:
        return None
    return value


def _load(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def checked_artifact(run: Path, relative: str, *, optional=False) -> Path | None:
    run = Path(run).resolve()
    manifest = _load(run / "run_manifest.json")
    if manifest.get("status") != "complete":
        raise ValueError("Evidence requires a completed run")
    path = (run / relative).resolve()
    if not path.is_relative_to(run):
        raise ValueError("Artifact path escapes completed run")
    expected = manifest.get("artifact_sha256", {}).get(relative)
    if expected is None and optional and not path.exists():
        return None
    if not expected or not path.is_file() or sha256_file(path) != expected:
        raise ValueError(f"Completed run artifact missing or modified: {relative}")
    return path


def checked_table(run: Path, name: str, *, optional=False) -> pd.DataFrame:
    import pyarrow.parquet as pq
    path = checked_artifact(run, f"gold/{name}.parquet", optional=optional)
    if path is None:
        return pd.DataFrame()
    if pq.ParquetFile(path).metadata.num_rows > MAX_TABLE_ROWS:
        raise ValueError(f"Workbench table row limit exceeded: {name}")
    return pd.read_parquet(path)


def identity_rows(frame: pd.DataFrame, row: dict, *, test_type=False) -> pd.DataFrame:
    if frame.empty:
        return frame.copy()
    required = {"cell_id", "test_id"} | ({"test_type"} if test_type else set())
    if not required.issubset(frame):
        raise ValueError(f"Evidence table lacks identity columns: {sorted(required - set(frame))}")
    mask = frame.cell_id.astype(str).eq(str(row.get("source_cell_id", row["cell_id"]))) & frame.test_id.astype(str).eq(str(row["test_id"]))
    if test_type:
        mask &= frame.test_type.astype(str).eq(str(row["test_type"]))
    if "include" in frame:
        mask &= frame["include"].eq(True)
    return frame.loc[mask].copy()


def mapped_cycles(run: Path, row: dict) -> pd.DataFrame:
    cycles = identity_rows(checked_table(run, "cycles"), row, test_type=True)
    if not {"cycle", "life_index"}.issubset(cycles):
        raise ValueError("No cycle/life_index mapping")
    cycles["cycle"] = pd.to_numeric(cycles.cycle, errors="coerce")
    if row.get("axis_kind") == "rpt_internal_cycle":
        if str(row.get("test_type", "")).strip().lower() != "rpt":
            raise ValueError("RPT internal cycle axis requires an RPT test")
        # Adapter-only view. The immutable gold life_index remains missing for
        # RPT; this explicit comparison coordinate does not claim lifetime age.
        cycles["life_index"] = cycles.cycle
    else:
        cycles["life_index"] = pd.to_numeric(cycles.life_index, errors="coerce")
    cycles = cycles.loc[np.isfinite(cycles.life_index)].copy()
    if not np.isfinite(cycles.cycle).all():
        raise ValueError("Cycle mapping contains nonfinite cycle identity")
    if cycles.cycle.duplicated().any() or cycles.life_index.duplicated().any():
        raise ValueError("Ambiguous cycle/life_index mapping")
    return cycles


def silver_batches(run: Path, row: dict, *, cycle=None, columns=None, required_signals=None):
    """Verified column batches; cycle row-group pruning and hard scan bounds."""
    import pyarrow.parquet as pq
    manifest = _load(Path(run) / "run_manifest.json")
    prefix = hashlib.sha256(str(row["test_id"]).encode()).hexdigest()[:16]
    candidates = sorted(p for p in manifest.get("artifact_sha256", {})
                        if p.startswith(f"silver/{prefix}/") and p.endswith(".parquet"))
    if not candidates:
        candidates = sorted(p for p in manifest.get("artifact_sha256", {})
                            if p.startswith("silver/") and p.endswith(".parquet"))
    if len(candidates) > 512:
        raise ValueError("Evidence partition limit exceeded")
    scanned = 0
    for relative in candidates:
        path = (Path(run) / relative).resolve()
        if not path.is_relative_to(Path(run).resolve()):
            raise ValueError("Silver path escapes completed run")
        checked_artifact(run, relative)
        parquet = pq.ParquetFile(path)
        names = parquet.schema_arrow.names
        required = {"cycle", "cell_id", "test_id"}
        if not required.issubset(names):
            raise ValueError("Normalized evidence lacks identity columns")
        if len(candidates) == 1 and required_signals is not None and not set(required_signals).issubset(names):
            # The artifact was still verified above. Do not decode a batch just
            # to discover that a deliberately omitted signal is unavailable.
            # For mixed legacy partitions, preserve row-level identity filtering
            # before deciding which signals the requested test actually has.
            return
        groups = []
        cycle_index = names.index("cycle")
        for number in range(parquet.num_row_groups):
            statistics = parquet.metadata.row_group(number).column(cycle_index).statistics
            if cycle is not None and statistics is not None and statistics.has_min_max:
                if not statistics.min <= cycle <= statistics.max:
                    continue
            groups.append(number)
        if not groups:
            continue
        chosen = None if columns is None else sorted((set(columns) | required | {"test_type"}) & set(names))
        for batch in parquet.iter_batches(batch_size=32768, row_groups=groups, columns=chosen):
            scanned += batch.num_rows
            if scanned > MAX_SCAN_ROWS:
                raise ValueError("Evidence scan row limit exceeded; narrow or repartition the input")
            # Native DB writers may retain pandas Arrow-extension metadata.
            # The analytical helpers expect NumPy-backed comparisons/cumsums;
            # physical Arrow field types remain intact when metadata is ignored.
            frame = identity_rows(batch.to_pandas(ignore_metadata=True), row, test_type=True)
            if cycle is not None:
                frame = frame.loc[frame.cycle.eq(cycle)]
            if not frame.empty:
                yield frame


def _upload_text(value, *, number_format=""):
    if value is None:
        return ""
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError("Registration cells must be finite; NaN/Infinity are not accepted")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if re.fullmatch(r"0+", number_format) and float(value).is_integer():
            return str(int(value)).zfill(len(number_format))
        return str(int(value)) if float(value).is_integer() else str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value)
    if text.lstrip().startswith("="):
        raise ValueError("Formulas are not accepted in registration uploads")
    if text.strip().lower() in {"nan", "inf", "+inf", "-inf", "infinity", "+infinity", "-infinity"}:
        raise ValueError("Registration cells must be finite")
    return text


def parse_registration_upload(content: bytes, filename: str, sheet: str | None = None) -> dict:
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("Registration upload is empty or exceeds 10 MiB")
    suffix = Path(filename).suffix.lower()
    sheets, selected = [], None
    if suffix == ".csv":
        try:
            decoded = content.decode("utf-8-sig")
            encoding = "utf-8-sig"
        except UnicodeDecodeError:
            decoded, encoding = content.decode("cp949"), "cp949"
        source = csv.reader(io.StringIO(decoded, newline=""), strict=True)
        raw_rows = []
        for values in source:
            if len(raw_rows) > MAX_UPLOAD_ROWS:
                raise ValueError("Registration upload row limit exceeded")
            if len(values) > MAX_UPLOAD_COLUMNS:
                raise ValueError("Registration upload column limit exceeded")
            raw_rows.append([_upload_text(value) for value in values])
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 50 MiB")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            sheets = list(book.sheetnames)
            if sheet is None and len(sheets) != 1:
                return {"columns": [], "rows": [], "sheets": sheets, "selected_sheet": None, "requires_sheet_selection": True}
            selected = sheet or sheets[0]
            if selected not in sheets:
                raise ValueError("Selected registration worksheet does not exist")
            worksheet = book[selected]
            if (worksheet.max_row or 0) > MAX_UPLOAD_ROWS + 1 or (worksheet.max_column or 0) > MAX_UPLOAD_COLUMNS:
                raise ValueError("Registration worksheet row/column limit exceeded")
            raw_rows = []
            for cells in worksheet.iter_rows():
                if len(raw_rows) > MAX_UPLOAD_ROWS or len(cells) > MAX_UPLOAD_COLUMNS:
                    raise ValueError("Registration worksheet row/column limit exceeded")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Formulas are not accepted in registration uploads")
                raw_rows.append([_upload_text(cell.value, number_format=cell.number_format) for cell in cells])
            encoding = "xlsx"
        finally:
            book.close()
    else:
        raise ValueError("Registration upload must be CSV or XLSX")
    if not raw_rows:
        raise ValueError("Registration upload has no header")
    columns = [value.strip() for value in raw_rows[0]]
    if not columns or any(not c for c in columns) or len(set(columns)) != len(columns):
        raise ValueError("Registration headers must be nonempty and unique")
    records = []
    for number, values in enumerate(raw_rows[1:], start=2):
        if not any(value.strip() for value in values):
            continue
        if len(values) != len(columns):
            raise ValueError(f"Registration row {number} has a different column count")
        records.append(dict(zip(columns, values)))
    if len(records) > MAX_UPLOAD_ROWS:
        raise ValueError("Registration upload row limit exceeded")
    return {"columns": columns, "rows": records, "sheets": sheets, "selected_sheet": selected,
            "requires_sheet_selection": False, "encoding": encoding}


def _full_external_observations(run, row, specs, *, cycles=None):
    if cycles is None:
        cycles = mapped_cycles(run, row)
    external = identity_rows(checked_table(run, "external_features", optional=True), row)
    if external.empty:
        return [], []
    mapping = dict(zip(cycles.cycle, cycles.life_index))
    observations, issues = [], []
    for spec in specs:
        selected = external.loc[external.metric_id.eq(spec["column"])].copy()
        if selected.empty:
            issues.append({"experiment_id": row["experiment_id"], "metric_id": spec["metric_id"], "reason": "external_metric_unavailable"})
            continue
        if selected.cycle.duplicated().any():
            raise ValueError("Duplicate external metric for one cycle")
        for key in ("unit", "definition_version", "source_revision"):
            if key not in selected or not selected[key].eq(spec[key]).all():
                raise ValueError(f"External metric {key} mismatch")
        if "namespace" in spec and ("namespace" not in selected or not selected.namespace.eq(spec["namespace"]).all()):
            raise ValueError("External metric namespace mismatch")
        if "role" in selected and not selected.role.eq("indicator").all():
            raise ValueError("External targets/QC cannot enter the workbench metric pool")
        for point in selected.to_dict("records"):
            life = mapping.get(point["cycle"])
            until = mapping.get(point.get("observed_until_cycle"))
            if life is None or until is None or until < life:
                issues.append({"experiment_id": row["experiment_id"], "metric_id": spec["metric_id"],
                               "reason": "external_observation_boundary_unmapped", "cycle": point["cycle"]})
                continue
            observations.append({"experiment_id": row["experiment_id"], "metric_id": spec["metric_id"],
                                 "life_index": float(life), "observed_until_life": float(until),
                                 "value": point["value"], "unit": spec["unit"]})
    return observations, issues


def _cached_run_extras(output, run, row, external_specs, *, cycles=None, code_hashes=None):
    """Cache verified derived values for immutable completed pipeline runs.

    Reuse checks manifest/source-code identity, artifact size/mtime_ns, and a
    separate cache SHA sidecar. Any changed artifact is reverified when read.
    This is a local performance cache, not protection from a party deliberately
    altering both bytes and filesystem timestamps. Evidence GET always hashes.
    """
    from . import metric_packs
    run = Path(run).resolve()
    manifest_path = run / "run_manifest.json"
    manifest = _load(manifest_path)
    if manifest.get("status") != "complete":
        raise ValueError("Workbench extras require a completed run")
    fingerprint = {}
    for relative in manifest.get("artifact_sha256", {}):
        artifact = (run / relative).resolve()
        if not artifact.is_relative_to(run):
            raise ValueError("Cached artifact path escapes completed run")
        stat = artifact.stat()
        fingerprint[relative] = [stat.st_size, stat.st_mtime_ns]
    if code_hashes is None:
        code_hashes = [sha256_file(Path(__file__)), sha256_file(Path(metric_packs.__file__))]
    key_data = {"version": 1, "run": str(run), "manifest": sha256_file(manifest_path),
                "code": code_hashes,
                "identity": {k: row.get(k) for k in ("experiment_id", "cell_id", "source_cell_id", "test_id", "test_type", "axis_kind")},
                "external_specs": external_specs}
    key = hashlib.sha256(json.dumps(key_data, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    cache_dir = Path(output) / "workbench_cache"
    cache_path, digest_path = cache_dir / f"{key}.json", cache_dir / f"{key}.sha256"
    try:
        if cache_path.is_file() and digest_path.is_file() and sha256_file(cache_path) == digest_path.read_text(encoding="ascii").strip():
            cached = _load(cache_path)
            if cached["artifact_stat"] == fingerprint:
                return cached["data"]
    except (ValueError, KeyError, OSError):
        pass
    if cycles is None and external_specs:
        cycles = mapped_cycles(run, row)
    observations, specs, issues = metric_packs.extract_metric_packs(run, row, cycles=cycles, source_revision=code_hashes[1])
    if external_specs:
        extra, extra_issues = _full_external_observations(run, row, external_specs, cycles=cycles)
        observations.extend(extra)
        issues.extend(extra_issues)
    result = json_safe({"observations": observations, "specs": specs, "issues": issues})
    serialized = json.dumps({"artifact_stat": fingerprint, "data": result}, ensure_ascii=False, allow_nan=False).encode()
    cache_dir.mkdir(parents=True, exist_ok=True)
    # Individual replacements are atomic; an interrupted pair is a cache miss.
    for destination, content in ((cache_path, serialized), (digest_path, hashlib.sha256(serialized).hexdigest().encode("ascii"))):
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(dir=cache_dir, delete=False) as stream:
                temporary = Path(stream.name)
                stream.write(content)
            os.replace(temporary, destination)
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)
    return result


def build_workbench_bundle(output: Path, config: dict, registrations: list[dict], payload: dict, *, observation_records=None, run_extras=None) -> dict:
    from . import metric_packs
    from .metric_packs import PACKS
    output = Path(output)
    experiments = [dict(row) for row in registrations if row.get("include", True) in (True, "true", "True", "1", 1)]
    by_id = {row["experiment_id"]: row for row in experiments}
    if len(by_id) != len(experiments):
        raise ValueError("Duplicate workbench experiment identity")
    obs_path = output / "observations.parquet"
    if observation_records is None and obs_path.exists():
        import pyarrow.parquet as pq
        if pq.ParquetFile(obs_path).metadata.num_rows > MAX_TABLE_ROWS:
            raise ValueError("Workbench observation row limit exceeded")
    observations = (pd.DataFrame(observation_records) if observation_records is not None else
                    pd.read_parquet(obs_path) if obs_path.exists() else pd.DataFrame())
    external_specs = [spec for spec in config["metrics"] if spec.get("table") == "external_features"]
    external_ids = {spec["metric_id"] for spec in external_specs}
    if not observations.empty:
        observations = observations.loc[~observations.metric_id.isin(external_ids)].copy()
        if "observed_until_life" not in observations:
            observations["observed_until_life"] = observations.life_index
    records = observations.to_dict("records")
    specs = {spec["metric_id"]: dict(spec, selected_by_default=spec.get("selected_by_default", True)) for spec in config["metrics"]}
    issues = []
    packs = {key: {"pack_id": key, **value, "selected_by_default": False, "metric_ids": [], "available": False} for key, value in PACKS.items()}
    # These are request-local inputs only. Completed artifact verification and
    # cache invalidation still run on every refresh; no global stat/hash cache.
    code_hashes = ([sha256_file(Path(__file__)), sha256_file(Path(metric_packs.__file__))]
                   if run_extras is None else None)
    for record in payload.get("ingestion", {}).get("records", []):
        row = by_id.get(record["experiment_id"])
        if row is None:
            continue
        run = Path(record["run_path"])
        extras = (run_extras[row["experiment_id"]] if run_extras is not None else
                  _cached_run_extras(output, run, row, external_specs, code_hashes=code_hashes))
        records.extend(extras["observations"])
        if len(records) > MAX_TABLE_ROWS:
            raise ValueError("Workbench observation row limit exceeded")
        issues.extend(extras["issues"])
        for spec in extras["specs"]:
            if spec["metric_id"] in specs and specs[spec["metric_id"]].get("family") != spec["family"]:
                raise ValueError("Conflicting optional metric definition")
            specs.setdefault(spec["metric_id"], dict(spec, selected_by_default=False, available=True))
            pack = packs.setdefault(spec["pack_id"], {"pack_id": spec["pack_id"], "selected_by_default": False, "metric_ids": []})
            if spec["metric_id"] not in pack["metric_ids"]:
                pack["metric_ids"].append(spec["metric_id"])
            pack["available"] = True
    # Temporary pending/failed raw runs must not erase a user's metric choices.
    # Retain known definitions only; observations always come from this update.
    previous_path = output / "dashboard" / "data.json"
    if previous_path.is_file():
        previous = _load(previous_path)
        for spec in previous.get("metrics", []):
            if (spec.get("table") != "workbench_pack" or spec.get("role") != "indicator"
                    or not str(spec.get("metric_id", "")).startswith("pack.") or spec.get("pack_id") not in PACKS):
                continue
            identity = spec["metric_id"]
            if identity in specs:
                continue
            required = ("column", "unit", "family", "definition_version", "source_revision")
            if any(not isinstance(spec.get(key), str) or not spec[key] for key in required):
                raise ValueError(f"Previously published optional metric definition is incomplete: {identity}")
            specs[identity] = dict(spec, selected_by_default=False, available=False,
                                   availability_reason="current_data_unavailable")
            issues.append({"metric_id": identity, "pack_id": spec["pack_id"],
                           "reason": "known_optional_metric_currently_unavailable"})
            packs[spec["pack_id"]]["metric_ids"].append(identity)
    for pack in packs.values():
        pack["available_metric_ids"] = [identity for identity in pack["metric_ids"] if specs[identity].get("available", False)]
        pack["unavailable_metric_ids"] = [identity for identity in pack["metric_ids"] if not specs[identity].get("available", False)]
        pack["available"] = bool(pack["available_metric_ids"])
    identities = [(r["experiment_id"], r["metric_id"], r["life_index"]) for r in records]
    if len(set(identities)) != len(identities):
        raise ValueError("Duplicate workbench observation identity")
    return json_safe({"schema_version": "workbench_bundle_v1", "experiments": experiments,
                      "observations": records, "metrics": list(specs.values()), "design_fields": config.get("design", {}).get("fields", []),
                      "control_fields": config.get("workbench", {}).get("control_fields", CONTROL_FIELDS),
                      "comparison": config["comparison"], "design_settings": config.get("design", {}).get("settings", {}),
                      "base_payload": payload, "generated_at": datetime.now(timezone.utc).isoformat(),
                      "is_demo": bool(config.get("is_demo", False)), "available_packs": list(packs.values()), "issues": issues})


def _display_derivatives(segment, direction, settings):
    empty = ([], [], {"status": "not_supported", "reason": "requires_valid_monotonic_CC_segment"})
    if direction not in {"charge", "discharge"} or not segment["mode"].fillna("").str.upper().eq("CC").all():
        return empty
    if _curve_features(segment, direction, settings)["curve_points"] < settings["min_curve_points"]:
        return empty
    t, v, current = [pd.to_numeric(segment[c], errors="coerce").to_numpy(float) for c in ("time_s", "voltage_v", "current_a")]
    q = np.r_[0., np.cumsum((np.abs(current[:-1]) + np.abs(current[1:])) / 2 * np.diff(t) / 3600)]
    valid = (1 if direction == "charge" else -1) * np.diff(v) > settings["min_voltage_delta_v"]
    starts, ends = np.r_[0, np.flatnonzero(~valid) + 1], np.r_[np.flatnonzero(~valid) + 1, len(v)]
    a, b = max(((a, b) for a, b in zip(starts, ends) if b - a >= settings["min_curve_points"]), key=lambda pair: pair[1] - pair[0])
    vv, qq = v[a:b], q[a:b]
    dq, dv = np.diff(qq), np.diff(vv)
    dq_dv, dv_dq = np.abs(dq / dv), np.abs(dv / dq)
    v_mid, q_mid = (vv[:-1] + vv[1:]) / 2, (qq[:-1] + qq[1:]) / 2
    if len(dq_dv) >= 5:
        dq_dv = np.convolve(dq_dv, np.ones(5) / 5, mode="valid")
        dv_dq = np.convolve(dv_dq, np.ones(5) / 5, mode="valid")
        v_mid, q_mid = v_mid[2:-2], q_mid[2:-2]
    indices = np.unique(np.linspace(0, len(dq_dv) - 1, min(MAX_CURVE_POINTS, len(dq_dv)), dtype=int))
    return ([{"voltage_v": float(v_mid[i]), "value": float(dq_dv[i])} for i in indices],
            [{"capacity_ah": float(q_mid[i]), "value": float(dv_dq[i])} for i in indices],
            {"status": "valid", "method": "chronological longest monotonic CC run; absolute derivatives; five-point moving average",
             "capacity_basis": "integrated_abs_current_Ah_from_segment_start", "dqdv_unit": "Ah/V", "dvdq_unit": "V/Ah",
             "source_row_start": segment.iloc[a].get("source_row"), "source_row_end": segment.iloc[b - 1].get("source_row"),
             "downsampled": len(indices) < len(dq_dv), "original_derivative_points": len(dq_dv)})


def read_evidence(output: Path, experiment_id: str, cycle: float | None = None, *, reference=None) -> dict:
    output = Path(output).resolve()
    if reference is not None:
        from .database_workbench import checked_database_reference
        row, run, source, snapshot = checked_database_reference(reference)
        if row["experiment_id"] != experiment_id:
            raise ValueError("Database evidence reference has a different experiment identity")
        entry = {"raw_sha256": reference["source_sha256"]}
        provenance = [source]
    else:
        state = _load(output / "update_state.json")
        entry = state.get("records", {}).get(experiment_id)
        if entry is None:
            raise ValueError("Experiment has no completed raw run; it may be pending")
        registry = pd.read_csv(output / "experiments.csv", dtype=str, keep_default_na=False)
        match = registry.loc[registry.experiment_id.eq(experiment_id)]
        if len(match) != 1:
            raise ValueError("Experiment identity missing or ambiguous")
        row = match.iloc[0].to_dict()
        run = Path(entry["run_path"]).resolve()
        if not run.is_relative_to(output / "lake"):
            raise ValueError("Evidence run is outside this workbench lake")
        # Recover test_type only from the verified legacy source mapping.
        snapshot = _load(checked_artifact(run, "config_snapshot.json"))
        sources = [s for s in snapshot.get("sources", []) if str(s.get("cell_id")) == row["cell_id"] and str(s.get("test_id")) == row["test_id"]]
        if len(sources) != 1:
            raise ValueError("Verified source mapping is missing or ambiguous")
        row["test_type"] = sources[0]["test_type"]
        inputs = _load(checked_artifact(run, "input_snapshot.json"))
        provenance = [item for item in inputs if item.get("kind") == "raw" and str(item.get("cell_id")) == row["cell_id"] and str(item.get("test_id")) == row["test_id"]]
        if len(provenance) != 1 or provenance[0].get("sha256") != entry.get("raw_sha256"):
            raise ValueError("Evidence raw provenance differs from the completed registration run")
    cycles = mapped_cycles(run, row).sort_values("cycle")
    if cycles.empty:
        raise ValueError("No cycles are available for this experiment")
    if cycle is None:
        cycle = float(cycles.iloc[0].cycle)
    if not math.isfinite(float(cycle)) or not cycles.cycle.eq(float(cycle)).any():
        raise ValueError("Requested cycle is not available for this experiment")
    chunks, count = [], 0
    for batch in silver_batches(run, row, cycle=float(cycle)):
        count += len(batch)
        if count > MAX_CYCLE_POINTS:
            raise ValueError("Selected cycle exceeds evidence point limit")
        chunks.append(batch)
    frame = pd.concat(chunks, ignore_index=True) if chunks else pd.DataFrame()
    curves, warnings = [], []
    signal_names = ("time_s", "step_time_s", "voltage_v", "current_a", "capacity_ah", "soc_pct", "mode")
    available_signals = [name for name in signal_names if name in frame]
    unavailable_signals = [name for name in signal_names if name not in frame]
    settings = DEFAULTS | snapshot.get("features", {})
    if not frame.empty:
        if provenance[0].get("kind") == "pne_restore":
            members = {item.get("member", item.get("path")): item["sha256"] for item in provenance[0].get("members", [])}
            if not {"source_member", "source_member_row"}.issubset(frame) or not members:
                raise ValueError("Native evidence requires source member hashes and physical member rows")
            if not frame.source_hash.eq(frame.source_member.map(members)).all():
                raise ValueError("Normalized evidence member hash differs from run input snapshot")
        elif not frame.source_hash.eq(provenance[0]["sha256"]).all():
            raise ValueError("Normalized evidence source hash differs from run input snapshot")
        frame = frame.sort_values("source_row", kind="stable")
        if unavailable_signals:
            warnings.append({"reason": "signal_not_retained", "fields": unavailable_signals})
        for name, default in {"mode": "unknown", "current_a": np.nan, "soc_pct": np.nan, "capacity_ah": np.nan, "step_time_s": np.nan}.items():
            if name not in frame:
                frame[name] = default
        parts = _segments(frame)
        if len(parts) > 64:
            raise ValueError("Selected cycle exceeds evidence segment limit")
        for segment in parts:
            direction_basis = "measured_current"
            if "current_a" in available_signals:
                leg = _direction(segment, settings)
            else:
                direction_basis = "recorded_step_kind"
                kinds = set(segment.get("step_kind", pd.Series(dtype=str)).dropna().astype(str).str.strip().str.lower())
                leg = next(iter(kinds)) if len(kinds) == 1 and kinds <= {"charge", "discharge", "rest"} else "unknown"
            chosen = np.unique(np.linspace(0, len(segment) - 1, min(MAX_CURVE_POINTS, len(segment)), dtype=int))
            point_columns = ["time_s", "voltage_v", "current_a", "capacity_ah", "soc_pct", "source_row"]
            point_columns += [key for key in ("source_member", "source_member_row") if key in segment]
            points = segment.iloc[chosen][point_columns].to_dict("records")
            if "current_a" not in available_signals or "mode" not in available_signals:
                dq, dv, derivative = [], [], {"status": "not_supported", "reason": "current_or_explicit_CC_mode_not_retained"}
            else:
                dq, dv, derivative = _display_derivatives(segment, leg, settings)
            first, last = segment.iloc[0], segment.iloc[-1]
            curves.append({"cycle": float(cycle), "leg": leg, "direction_basis": direction_basis, "segment_id": first.segment_id,
                           "step_id": first.step_id, "points": points, "dqdv": dq, "dvdq": dv, "derivative": derivative,
                           "source": {"path": first.source_path, "sha256": first.source_hash,
                                      "row_start": first.source_row, "row_end": last.source_row,
                                      "row_basis": row.get("source_row_basis", "1-based data record; not physical file line"),
                                      "member": first.get("source_member"), "member_row_start": first.get("source_member_row"),
                                      "last_member": last.get("source_member"), "member_row_end": last.get("source_member_row")},
                           "sampling": {"original_points": len(segment), "returned_points": len(chosen),
                                        "downsampled": len(chosen) < len(segment), "method": "evenly spaced source rows including endpoints"}})
    else:
        warnings.append({"reason": "normalized_curve_not_available"})
    tables = {}
    for name in ("pulses", "diagnostics", "quality"):
        table = identity_rows(checked_table(run, name, optional=True), row)
        if not table.empty:
            table = table.loc[pd.to_numeric(table.cycle, errors="coerce").eq(cycle)]
            if len(table) > 10000:
                raise ValueError(f"Evidence row limit exceeded: {name}")
        tables[name] = table.to_dict("records")
    if reference is not None:
        checked_database_reference(reference)  # Reject a concurrently replaced run manifest.
    return json_safe({"schema_version": "experiment_evidence_v1", "experiment_id": experiment_id, "selected_cycle": float(cycle),
                      "cycles": cycles[["cycle", "life_index"]].to_dict("records"), "curves": curves, **tables,
                      "warnings": warnings, "source_hash": entry.get("raw_sha256"), "run_path": str(run),
                      "axis_kind": row.get("axis_kind", "life_index"), "axis_label": row.get("axis_label", "Life index"),
                      "available_signals": available_signals, "unavailable_signals": unavailable_signals,
                      "storage_retention": row.get("retention"),
                      "provenance_scope": "Completed run snapshot; original raw file is not reread or modified by evidence retrieval",
                      "limits": {"max_cycle_points": MAX_CYCLE_POINTS, "max_points_per_curve": MAX_CURVE_POINTS,
                                 "max_scan_rows": MAX_SCAN_ROWS}, "raw_modified": False})
